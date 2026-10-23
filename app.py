import os, secrets, hashlib, hmac, time, uuid, io, json, re
from datetime import datetime, timezone
from pathlib import Path
from contextlib import asynccontextmanager
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Response, Request
from fastapi.responses import FileResponse
from sqlalchemy import create_engine, text
from sqlalchemy.exc import IntegrityError
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
from pydantic import BaseModel, Field
from openpyxl import load_workbook
import xlrd
from sqlalchemy import inspect

ROOT=Path(__file__).parent
URL=os.getenv('DATABASE_URL','sqlite:///'+str(ROOT/'demo.sqlite3'))
engine=create_engine(URL, pool_pre_ping=True, connect_args={'check_same_thread':False} if URL.startswith('sqlite') else {})
SECRET=os.getenv('APP_SECRET','')
if not SECRET or len(SECRET)<32: raise RuntimeError('APP_SECRET must be at least 32 characters')
signer=URLSafeTimedSerializer(SECRET,salt='wms-session-v1')
SCHEMA=[
'CREATE TABLE IF NOT EXISTS employee_permissions(user_id VARCHAR(36) NOT NULL,permission VARCHAR(50) NOT NULL,enabled INTEGER NOT NULL,PRIMARY KEY(user_id,permission))',
'CREATE TABLE IF NOT EXISTS employee_permission_versions(user_id VARCHAR(36) PRIMARY KEY,revision INTEGER NOT NULL)',
'CREATE TABLE IF NOT EXISTS forecast_original_files(id VARCHAR(36) PRIMARY KEY,filename VARCHAR(255) NOT NULL,content_base64 TEXT NOT NULL,sha256 VARCHAR(64) NOT NULL,user_id VARCHAR(36) NOT NULL,customer_id VARCHAR(36),container_id VARCHAR(36),created_at VARCHAR(40) NOT NULL)',
'CREATE TABLE IF NOT EXISTS piece_labels(code VARCHAR(100) PRIMARY KEY, order_id VARCHAR(36) NOT NULL, piece_no INTEGER NOT NULL, operation VARCHAR(30) NOT NULL, user_id VARCHAR(36) NOT NULL, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)',
'CREATE TABLE IF NOT EXISTS wms_brand_migrations(id VARCHAR(100) PRIMARY KEY, detail TEXT NOT NULL)',
'CREATE TABLE IF NOT EXISTS container_import_receipts(request_id VARCHAR(100) PRIMARY KEY, user_id VARCHAR(36) NOT NULL, payload_hash VARCHAR(64) NOT NULL, response TEXT NOT NULL)',
"CREATE TABLE IF NOT EXISTS client_forecasts(id VARCHAR(36) PRIMARY KEY, customer_id VARCHAR(36) NOT NULL, warehouse_id VARCHAR(36) NOT NULL, code VARCHAR(120) NOT NULL DEFAULT '', status VARCHAR(20) NOT NULL DEFAULT 'draft', filename VARCHAR(255) NOT NULL DEFAULT '', rows_json TEXT NOT NULL, revision INTEGER NOT NULL DEFAULT 1, container_id VARCHAR(36), created_by VARCHAR(36) NOT NULL, submitted_by VARCHAR(36), created_at VARCHAR(40) NOT NULL, updated_at VARCHAR(40) NOT NULL, submitted_at VARCHAR(40))",
'''CREATE TABLE IF NOT EXISTS warehouses(id VARCHAR(36) PRIMARY KEY, name VARCHAR(120) NOT NULL, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS users(id VARCHAR(36) PRIMARY KEY, username VARCHAR(100) UNIQUE NOT NULL, password_hash TEXT NOT NULL, role VARCHAR(30) NOT NULL, active INTEGER NOT NULL DEFAULT 1, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS memberships(user_id VARCHAR(36) NOT NULL, warehouse_id VARCHAR(36) NOT NULL, PRIMARY KEY(user_id,warehouse_id))''',
'''CREATE TABLE IF NOT EXISTS containers(id VARCHAR(36) PRIMARY KEY, warehouse_id VARCHAR(36) NOT NULL, code VARCHAR(120) NOT NULL, status VARCHAR(30) NOT NULL DEFAULT '预报', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, UNIQUE(warehouse_id,code))''',
'''CREATE TABLE IF NOT EXISTS orders(id VARCHAR(36) PRIMARY KEY, container_id VARCHAR(36) NOT NULL, system_no VARCHAR(150) NOT NULL, main_barcode VARCHAR(200), expected INTEGER NOT NULL, received INTEGER NOT NULL DEFAULT 0, mode VARCHAR(20) NOT NULL DEFAULT '待确认', UNIQUE(container_id,system_no))''',
'''CREATE TABLE IF NOT EXISTS barcodes(id VARCHAR(36) PRIMARY KEY, order_id VARCHAR(36) NOT NULL, barcode VARCHAR(200) NOT NULL, UNIQUE(order_id,barcode))''',
'''CREATE TABLE IF NOT EXISTS scans(id VARCHAR(36) PRIMARY KEY, request_id VARCHAR(100) UNIQUE NOT NULL, warehouse_id VARCHAR(36) NOT NULL, container_id VARCHAR(36) NOT NULL, order_id VARCHAR(36), barcode VARCHAR(200) NOT NULL, user_id VARCHAR(36) NOT NULL, result VARCHAR(30) NOT NULL, reason TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS exceptions(id VARCHAR(36) PRIMARY KEY, scan_id VARCHAR(36) UNIQUE NOT NULL, warehouse_id VARCHAR(36) NOT NULL, container_id VARCHAR(36) NOT NULL, barcode VARCHAR(200) NOT NULL, reason TEXT, status VARCHAR(30) DEFAULT '待处理', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS audit_logs(id VARCHAR(36) PRIMARY KEY, warehouse_id VARCHAR(36), user_id VARCHAR(36) NOT NULL, action VARCHAR(100) NOT NULL, detail TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS customers(id VARCHAR(36) PRIMARY KEY, name VARCHAR(120) NOT NULL UNIQUE, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS goods_notes(id VARCHAR(36) PRIMARY KEY, order_id VARCHAR(36) NOT NULL, user_id VARCHAR(36) NOT NULL, note TEXT NOT NULL, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS goods_photos(id VARCHAR(36) PRIMARY KEY, order_id VARCHAR(36) NOT NULL, user_id VARCHAR(36) NOT NULL, filename VARCHAR(200) NOT NULL, mime VARCHAR(50) NOT NULL, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS goods_events(id VARCHAR(36) PRIMARY KEY, order_id VARCHAR(36) NOT NULL, kind VARCHAR(30) NOT NULL, event_time VARCHAR(40) NOT NULL, user_id VARCHAR(36) NOT NULL, reason TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS order_pieces(id VARCHAR(36) PRIMARY KEY, order_id VARCHAR(36) NOT NULL, piece_no INTEGER NOT NULL, weight FLOAT, length FLOAT, width FLOAT, height FLOAT, note TEXT, source_row INTEGER, UNIQUE(order_id,piece_no))''',
'''CREATE TABLE IF NOT EXISTS status_history(id VARCHAR(36) PRIMARY KEY, order_id VARCHAR(36) NOT NULL, user_id VARCHAR(36) NOT NULL, old_status VARCHAR(30), new_status VARCHAR(30) NOT NULL, reason TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE TABLE IF NOT EXISTS piece_operations(id VARCHAR(36) PRIMARY KEY, request_id VARCHAR(100) UNIQUE NOT NULL, order_id VARCHAR(36) NOT NULL, warehouse_id VARCHAR(36) NOT NULL, user_id VARCHAR(36) NOT NULL, barcode VARCHAR(200) NOT NULL, operation VARCHAR(30) NOT NULL, result VARCHAR(30) NOT NULL, reason TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''',
'''CREATE UNIQUE INDEX IF NOT EXISTS ux_unique_package ON scans(container_id,barcode) WHERE result='成功-独立' '''
]
def uid(): return str(uuid.uuid4())
def migrate_return_checks(c):
 # Retain older tables, including their required scope column and history.
 columns={x['name'] for x in inspect(c).get_columns('return_checks')}
 for name,typ in {'scope':"VARCHAR(20) NOT NULL DEFAULT 'single'",'reason':"TEXT NOT NULL DEFAULT ''",'repaired_at':'VARCHAR(40)','repaired_by':'VARCHAR(36)'}.items():
  if name not in columns:c.execute(text(f'ALTER TABLE return_checks ADD COLUMN {name} {typ}'))
def ensure_damage_repaired(c,o):
 if c.execute(text("SELECT id FROM order_pieces WHERE order_id=:o AND damage_status='有破损' LIMIT 1"),{'o':o['id']}).first():
  fail('入库破损尚未全部确认修复，禁止出库或重出；请由主管在货物详情确认修复',409)
 if c.execute(text("SELECT id FROM return_checks WHERE order_id=:o AND cycle=:cy AND damage_choice='有破损' AND repaired_at IS NULL LIMIT 1"),{'o':o['id'],'cy':int(o['return_cycle'] or 0)}).first():
  fail('本轮退回破损尚未全部确认修复，禁止重出；请由主管在货物详情确认修复',409)
def migrate_warehouse_brand(c):
 key='aifuge-warehouse-76-v1'
 if c.execute(text('SELECT id FROM wms_brand_migrations WHERE id=:id'),{'id':key}).first():return
 candidates=c.execute(text("SELECT id,name FROM warehouses WHERE REPLACE(name,' ','') IN ('一号仓','快进快进一号仓')")).mappings().all()
 changes=[]
 if len(candidates)==1 and not c.execute(text("SELECT id FROM warehouses WHERE name='76号仓库'")).first():
  old=candidates[0]
  c.execute(text("UPDATE warehouses SET name='76号仓库' WHERE id=:id"),{'id':old['id']})
  changes.append({'id':old['id'],'old':old['name'],'new':'76号仓库'})
 c.execute(text('INSERT INTO wms_brand_migrations(id,detail) VALUES(:id,:d)'),{'id':key,'d':json.dumps(changes,ensure_ascii=False)})

def pw_hash(password):
 salt=secrets.token_bytes(16); key=hashlib.scrypt(password.encode(),salt=salt,n=16384,r=8,p=1)
 return salt.hex()+':'+key.hex()
def pw_ok(password,stored):
 try:
  s,k=stored.split(':'); return hmac.compare_digest(hashlib.scrypt(password.encode(),salt=bytes.fromhex(s),n=16384,r=8,p=1),bytes.fromhex(k))
 except (ValueError,TypeError):return False
@asynccontextmanager
async def lifespan(app):
 with engine.begin() as c:
  if URL.startswith('sqlite'):c.execute(text('PRAGMA journal_mode=WAL'))
  for sql in SCHEMA:c.execute(text(sql))
  migrate_warehouse_brand(c)
  # Additive migration: retain existing tables and data.
  for table,columns in {'client_forecasts':{'source_id':'VARCHAR(36)'},'customers':{'deleted_at':'VARCHAR(40)','contact':"VARCHAR(120) NOT NULL DEFAULT ''",'phone':"VARCHAR(80) NOT NULL DEFAULT ''"},'users':{'customer_id':'VARCHAR(36)','client_warehouse_id':'VARCHAR(36)'},'containers':{'customer_id':'VARCHAR(36)','remark':'TEXT','deleted_at':'VARCHAR(40)'},'orders':{'transfer_no':'VARCHAR(200)','remark':'TEXT','channel':'TEXT','country':'TEXT','city':'TEXT','postcode':'TEXT','weight':'FLOAT','length':'FLOAT','width':'FLOAT','height':'FLOAT','recipient':'TEXT','company':'TEXT','phone':'TEXT','email':'TEXT','address':'TEXT','declared_value':'FLOAT','description':'TEXT','received_override':'INTEGER','goods_status':'VARCHAR(30)','returned_count':'INTEGER NOT NULL DEFAULT 0','return_target':'INTEGER NOT NULL DEFAULT 0','reout_count':'INTEGER NOT NULL DEFAULT 0','return_cycle':'INTEGER NOT NULL DEFAULT 0','outbound_count':'INTEGER NOT NULL DEFAULT 0','hold_count':'INTEGER NOT NULL DEFAULT 0','hold_prealert':'INTEGER NOT NULL DEFAULT 0','hold_label':'INTEGER NOT NULL DEFAULT 0','damage_label':'INTEGER NOT NULL DEFAULT 0'}}.items():
   existing={x['name'] for x in inspect(c).get_columns(table)}
   for name,typ in columns.items():
    if name not in existing:c.execute(text(f'ALTER TABLE {table} ADD COLUMN {name} {typ}'))
  c.execute(text('CREATE TABLE IF NOT EXISTS return_checks(id VARCHAR(36) PRIMARY KEY, order_id VARCHAR(36) NOT NULL, cycle INTEGER NOT NULL, request_id VARCHAR(100) UNIQUE NOT NULL, quantity INTEGER NOT NULL, damage_choice VARCHAR(20) NOT NULL, user_id VARCHAR(36) NOT NULL, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)'))
  migrate_return_checks(c)
  c.execute(text('CREATE TABLE IF NOT EXISTS operation_receipts(request_id VARCHAR(100) PRIMARY KEY, user_id VARCHAR(36) NOT NULL, payload TEXT NOT NULL, response TEXT NOT NULL)'))
  # Additive migration for per-piece damage and repair status.
  piece_columns={x['name'] for x in inspect(c).get_columns('order_pieces')}
  if 'damage_status' not in piece_columns:c.execute(text('ALTER TABLE order_pieces ADD COLUMN damage_status VARCHAR(20)'))
  # Persist historical warning labels, including appended notes. Never clear them on edits.
  c.execute(text("UPDATE orders SET hold_label=1 WHERE COALESCE(remark,'') LIKE '%扣货%' OR EXISTS (SELECT 1 FROM goods_notes n WHERE n.order_id=orders.id AND n.note LIKE '%扣货%')"))
  c.execute(text("UPDATE orders SET damage_label=1 WHERE COALESCE(remark,'') LIKE '%破损%' OR EXISTS (SELECT 1 FROM goods_notes n WHERE n.order_id=orders.id AND n.note LIKE '%破损%')"))
  count=c.execute(text('SELECT COUNT(*) FROM users')).scalar()
  if count==0:
   user=os.getenv('BOOTSTRAP_ADMIN_USER','');password=os.getenv('BOOTSTRAP_ADMIN_PASSWORD','')
   if not user or len(password)<8:raise RuntimeError('Set BOOTSTRAP_ADMIN_USER and BOOTSTRAP_ADMIN_PASSWORD (8+ chars) for initial setup')
   wid=uid();admin=uid()
   c.execute(text('INSERT INTO warehouses(id,name) VALUES(:id,:name)'),{'id':wid,'name':os.getenv('INITIAL_WAREHOUSE_NAME','76号仓库')})
   c.execute(text('INSERT INTO users(id,username,password_hash,role) VALUES(:id,:u,:p,:r)'),{'id':admin,'u':user,'p':pw_hash(password),'r':'老板'})
   c.execute(text('INSERT INTO memberships(user_id,warehouse_id) VALUES(:u,:w)'),{'u':admin,'w':wid})
 yield
app=FastAPI(title='爱超凡 aifuge 仓储系统',lifespan=lifespan)
def fail(msg,status=400):raise HTTPException(status,msg)
def current(request:Request):
 token=request.cookies.get('wms_session')
 if not token:fail('请登录',401)
 try: user_id=signer.loads(token,max_age=12*3600)
 except (BadSignature,SignatureExpired):fail('登录已过期',401)
 with engine.connect() as c:
  u=c.execute(text('SELECT id,username,role,active,customer_id,client_warehouse_id FROM users WHERE id=:id'),{'id':user_id}).mappings().first()
 if not u or not u['active']:fail('账号不可用',401)
 if u['role']=='客户' and request.url.path not in ('/api/me','/api/logout') and not request.url.path.startswith('/api/client/'):
  fail('客户账号仅可访问本公司的客户端功能',403)
 result=dict(u)
 result['permissions']=sorted(effective_permissions(result))
 return result
def protect(request:Request,user=Depends(current)):
 # Cross-site form and cookie attacks: require matching origin for state-changing calls.
 origin=request.headers.get('origin');host=request.headers.get('host','')
 if origin and origin.rstrip('/').split('://')[-1]!=host:fail('来源校验失败',403)
 return user
def allow(user,roles):
 if user['role'] not in roles:fail('权限不足',403)
def access(c,user,wid):
 if user['role']=='老板':
  # owner can access only assigned warehouses, not unrelated organizations
  pass
 x=c.execute(text('SELECT 1 FROM memberships WHERE user_id=:u AND warehouse_id=:w'),{'u':user['id'],'w':wid}).first()
 if not x:fail('无权访问此仓库',403)
def audit(c,u,w,action,detail=''):
 c.execute(text('INSERT INTO audit_logs(id,warehouse_id,user_id,action,detail) VALUES(:id,:w,:u,:a,:d)'),{'id':uid(),'w':w,'u':u,'a':action,'d':detail})
class Login(BaseModel):username:str;password:str
@app.post('/api/login')
def login(x:Login,response:Response):
 with engine.connect() as c:u=c.execute(text('SELECT * FROM users WHERE username=:u'),{'u':x.username}).mappings().first()
 if not u or not u['active'] or not pw_ok(x.password,u['password_hash']):fail('账号或密码错误',401)
 response.set_cookie('wms_session',signer.dumps(u['id']),max_age=43200,httponly=True,secure=os.getenv('COOKIE_SECURE','true').lower()=='true',samesite='strict',path='/')
 return {'username':u['username'],'role':u['role']}
@app.post('/api/logout')
def logout(response:Response,user=Depends(protect)):
 response.delete_cookie('wms_session',path='/');return {'ok':True}
@app.get('/api/me')
def me(user=Depends(current)):return user
@app.get('/api/warehouses')
def warehouses(user=Depends(current)):
 with engine.connect() as c:return [dict(r) for r in c.execute(text('SELECT w.* FROM warehouses w JOIN memberships m ON m.warehouse_id=w.id WHERE m.user_id=:u ORDER BY w.name'),{'u':user['id']}).mappings()]
class NewWarehouse(BaseModel):name:str=Field(min_length=1,max_length=120)
@app.post('/api/warehouses')
def add_warehouse(x:NewWarehouse,user=Depends(protect)):
 permit(user,'warehouses');wid=uid()
 if not x.name.strip():fail('仓库名称不能为空')
 with engine.begin() as c:
  if c.execute(text('SELECT id FROM warehouses WHERE LOWER(name)=LOWER(:n)'),{'n':x.name.strip()}).first():fail('仓库名称已存在',409)
  c.execute(text('INSERT INTO warehouses(id,name) VALUES(:id,:n)'),{'id':wid,'n':x.name.strip()})
  c.execute(text('INSERT INTO memberships(user_id,warehouse_id) VALUES(:u,:w)'),{'u':user['id'],'w':wid})
  audit(c,user['id'],wid,'创建仓库',x.name)
 return {'id':wid}
@app.get('/api/users')
def users(user=Depends(current)):
 permit(user,'staff')
 with engine.connect() as c:return [dict(r) for r in c.execute(text("SELECT u.id,u.username,u.role,u.active,c.name customer_name FROM users u LEFT JOIN customers c ON c.id=u.customer_id WHERE (:boss=1 OR u.role NOT IN ('老板','客户')) ORDER BY u.username"),{'boss':int(user['role']=='老板')}).mappings()]
class NewUser(BaseModel):username:str=Field(min_length=2,max_length=100);password:str=Field(min_length=8);role:str;warehouse_ids:list[str]
@app.post('/api/users')
def create_user(x:NewUser,user=Depends(protect)):
 permit(user,'staff')
 if x.role not in ('老板','主管','客服','仓库管理员','PDA操作员','PDA员工','只读人员'):fail('角色不合法')
 if user['role']!='老板' and x.role=='老板':fail('主管不能创建老板账号',403)
 if not x.warehouse_ids:fail('至少分配一个仓库')
 with engine.begin() as c:
  for wid in set(x.warehouse_ids):access(c,user,wid)
  ident=uid()
  try:c.execute(text('INSERT INTO users(id,username,password_hash,role) VALUES(:id,:u,:p,:r)'),{'id':ident,'u':x.username,'p':pw_hash(x.password),'r':x.role})
  except IntegrityError:fail('账号已存在')
  for wid in set(x.warehouse_ids):c.execute(text('INSERT INTO memberships(user_id,warehouse_id) VALUES(:u,:w)'),{'u':ident,'w':wid})
  audit(c,user['id'],None,'创建员工',x.username+' '+x.role)
 return {'id':ident}
class UserUpdate(BaseModel):active:bool
@app.patch('/api/users/{target}')
def change_user(target:str,x:UserUpdate,user=Depends(protect)):
 permit(user,'staff')
 if target==user['id'] and not x.active:fail('不能停用自己的账号')
 with engine.begin() as c:
  target_role=c.execute(text('SELECT role FROM users WHERE id=:id'),{'id':target}).scalar()
  if target_role is None:fail('账号不存在',404)
  if user['role']!='老板' and target_role in ('老板','客户'):fail('主管不能修改该账号',403)
  c.execute(text('UPDATE users SET active=:a WHERE id=:id'),{'a':int(x.active),'id':target})
  audit(c,user['id'],None,'更改员工状态',target+':'+str(x.active))
 return {'ok':True}
class NewCustomer(BaseModel):name:str=Field(min_length=1,max_length=120)
@app.get('/api/customers')
def customers(user=Depends(current)):
 with engine.connect() as c:return [dict(x) for x in c.execute(text('SELECT * FROM customers WHERE deleted_at IS NULL ORDER BY name')).mappings()]
@app.post('/api/customers')
def add_customer(x:NewCustomer,user=Depends(protect)):
 permit(user,'customers')
 with engine.begin() as c:
  ident=uid()
  try:c.execute(text('INSERT INTO customers(id,name) VALUES(:id,:n)'),{'id':ident,'n':x.name.strip()})
  except IntegrityError:fail('客户名称已存在')
  audit(c,user['id'],None,'新增客户',x.name)
 return {'id':ident}
class NewContainer(BaseModel):warehouse_id:str;customer_id:str="";code:str=Field(min_length=1,max_length=120)
@app.post('/api/containers')
def add_container(x:NewContainer,user=Depends(protect)):
 permit(user,'forecast')
 with engine.begin() as c:
  access(c,user,x.warehouse_id);cid=uid()
  if x.customer_id and not c.execute(text('SELECT 1 FROM customers WHERE id=:id'),{'id':x.customer_id}).first():fail('客户不存在')
  try:c.execute(text('INSERT INTO containers(id,warehouse_id,code,customer_id) VALUES(:id,:w,:code,:customer)'),{'id':cid,'w':x.warehouse_id,'code':x.code.strip(),'customer':x.customer_id or None})
  except IntegrityError:fail('柜号已存在')
  audit(c,user['id'],x.warehouse_id,'创建货柜',x.code)
 return {'id':cid}
@app.get('/api/containers')
def list_containers(warehouse_id:str,q:str='',user=Depends(current)):
 permit(user,'read')
 with engine.connect() as c:
  access(c,user,warehouse_id)
  return [dict(r) for r in c.execute(text('''SELECT c.id,c.code,c.status,c.warehouse_id,c.customer_id,c.remark,cu.name customer_name,COALESCE(SUM(o.expected),0) expected,COALESCE(SUM(o.received),0) received FROM containers c LEFT JOIN orders o ON o.container_id=c.id LEFT JOIN customers cu ON cu.id=c.customer_id WHERE c.warehouse_id=:w AND c.deleted_at IS NULL AND (:q='' OR LOWER(c.code) LIKE :like OR LOWER(COALESCE(cu.name,'')) LIKE :like OR EXISTS(SELECT 1 FROM orders z WHERE z.container_id=c.id AND (LOWER(z.system_no) LIKE :like OR LOWER(COALESCE(z.main_barcode,'')) LIKE :like OR LOWER(COALESCE(z.transfer_no,'')) LIKE :like))) GROUP BY c.id,c.code,c.status,c.warehouse_id,c.customer_id,c.remark,cu.name ORDER BY c.created_at DESC'''),{'w':warehouse_id,'q':q,'like':'%'+q.lower()+'%'}).mappings()]
@app.get('/api/containers/{cid}')
def container_detail(cid:str,user=Depends(current)):
 permit(user,'read')
 with engine.connect() as c:
  ct=c.execute(text('SELECT * FROM containers WHERE id=:id'),{'id':cid}).mappings().first()
  if not ct or ct['deleted_at']:fail('货柜不存在',404)
  access(c,user,ct['warehouse_id'])
  orders=[dict(r) for r in c.execute(text('''SELECT o.*, (SELECT n.note FROM goods_notes n WHERE n.order_id=o.id ORDER BY n.created_at DESC LIMIT 1) latest_note, (SELECT COUNT(*) FROM goods_photos p WHERE p.order_id=o.id) photo_count, (SELECT COUNT(*) FROM order_pieces p WHERE p.order_id=o.id AND p.damage_status IN ('有破损','已修复')) damage_piece_count, (SELECT COUNT(*) FROM order_pieces p WHERE p.order_id=o.id AND p.damage_status='已修复') repaired_piece_count FROM orders o WHERE o.container_id=:id ORDER BY o.system_no'''),{'id':cid}).mappings()]
  scans=[dict(r) for r in c.execute(text('SELECT s.*,u.username FROM scans s JOIN users u ON u.id=s.user_id WHERE s.container_id=:id ORDER BY s.created_at DESC LIMIT 50'),{'id':cid}).mappings()]
  ex=[dict(r) for r in c.execute(text('SELECT * FROM exceptions WHERE container_id=:id ORDER BY created_at DESC LIMIT 50'),{'id':cid}).mappings()]
  piece_map={}
  for piece in c.execute(text('SELECT p.* FROM order_pieces p JOIN orders o ON o.id=p.order_id WHERE o.container_id=:id ORDER BY p.piece_no'),{'id':cid}).mappings():piece_map.setdefault(piece['order_id'],[]).append(dict(piece))
  for o in orders:o['pieces']=piece_map.get(o['id'],[])
  return {'container':dict(ct),'orders':[enriched(o) for o in orders],'scans':scans,'exceptions':ex}
class ContainerEdit(BaseModel):remark:str|None=None;code:str|None=None;customer_id:str|None=None
@app.patch('/api/containers/{cid}')
def edit_container(cid:str,x:ContainerEdit,user=Depends(protect)):
 permit(user,'edit')
 with engine.begin() as c:
  ct=c.execute(text('SELECT * FROM containers WHERE id=:id AND deleted_at IS NULL'),{'id':cid}).mappings().first()
  if not ct:fail('货柜不存在',404)
  access(c,user,ct['warehouse_id'])
  if x.customer_id is not None and x.customer_id!=ct['customer_id']:
   if c.execute(text('SELECT 1 FROM orders WHERE container_id=:id'),{'id':cid}).first():fail('已有预报，不可修改客户')
   if x.customer_id and not c.execute(text('SELECT 1 FROM customers WHERE id=:id'),{'id':x.customer_id}).first():fail('客户不存在')
   c.execute(text('UPDATE containers SET customer_id=:v WHERE id=:id'),{'v':x.customer_id or None,'id':cid})
  if x.remark is not None:c.execute(text('UPDATE containers SET remark=:v WHERE id=:id'),{'v':x.remark,'id':cid})
  if x.code is not None:
   if not x.code.strip():fail('柜号不能为空')
   c.execute(text('UPDATE containers SET code=:v WHERE id=:id'),{'v':x.code.strip(),'id':cid})
  audit(c,user['id'],ct['warehouse_id'],'修改货柜',json.dumps(x.model_dump(exclude_unset=True),ensure_ascii=False))
 return {'ok':True}
@app.delete('/api/containers/{cid}')
def delete_container(cid:str,user=Depends(protect)):
 allow(user,['老板','主管'])
 fail('请使用强制删除窗口，核对在库件数、完整柜号并填写原因',409)
class OrderEdit(BaseModel):
 system_no:str|None=None
 transfer_no:str|None=None
 main_barcode:str|None=None
 expected:int|None=Field(default=None,ge=0)
 received:int|None=Field(default=None,ge=0)
 remark:str|None=None
 mode:str|None=None
 channel:str|None=None
 country:str|None=None
 city:str|None=None
 postcode:str|None=None
 weight:float|None=None
 length:float|None=None
 width:float|None=None
 height:float|None=None
 recipient:str|None=None
 company:str|None=None
 phone:str|None=None
 email:str|None=None
 address:str|None=None
 declared_value:float|None=Field(default=None,ge=0,allow_inf_nan=False)
 description:str|None=None
@app.get('/api/orders/{oid}')
def order_detail(oid:str,user=Depends(current)):
 permit(user,'read')
 with engine.connect() as c:
  o=c.execute(text('SELECT o.*,ct.code container_code,ct.warehouse_id,cu.name customer_name FROM orders o JOIN containers ct ON ct.id=o.container_id LEFT JOIN customers cu ON cu.id=ct.customer_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('货物不存在',404)
  access(c,user,o['warehouse_id'])
  scans=[dict(x) for x in c.execute(text('SELECT s.*,u.username FROM scans s JOIN users u ON u.id=s.user_id WHERE s.order_id=:id ORDER BY s.created_at DESC'),{'id':oid}).mappings()]
  notes=[dict(x) for x in c.execute(text('SELECT n.*,u.username FROM goods_notes n JOIN users u ON u.id=n.user_id WHERE n.order_id=:id ORDER BY n.created_at DESC'),{'id':oid}).mappings()]
  photos=[dict(x) for x in c.execute(text('SELECT p.id,p.created_at,p.mime,u.username FROM goods_photos p JOIN users u ON u.id=p.user_id WHERE p.order_id=:id ORDER BY p.created_at DESC'),{'id':oid}).mappings()]
  events=[dict(x) for x in c.execute(text('SELECT e.*,u.username FROM goods_events e JOIN users u ON u.id=e.user_id WHERE e.order_id=:id ORDER BY e.created_at DESC'),{'id':oid}).mappings()]
  pieces=[dict(x) for x in c.execute(text('SELECT * FROM order_pieces WHERE order_id=:id ORDER BY piece_no'),{'id':oid}).mappings()]
  history=[dict(x) for x in c.execute(text('SELECT h.*,u.username FROM status_history h JOIN users u ON u.id=h.user_id WHERE h.order_id=:id ORDER BY h.created_at DESC,h.id DESC'),{'id':oid}).mappings()]
  return {'order':enriched(o),'return_checks':[dict(x) for x in c.execute(text('SELECT r.*,u.username FROM return_checks r JOIN users u ON u.id=r.user_id WHERE r.order_id=:id ORDER BY r.created_at DESC'),{'id':oid}).mappings()],'piece_operations':[dict(x) for x in c.execute(text('SELECT p.*,u.username FROM piece_operations p JOIN users u ON u.id=p.user_id WHERE p.order_id=:id ORDER BY p.created_at DESC'),{'id':oid}).mappings()],'pieces':pieces,'scans':scans,'notes':notes,'events':events,'photos':photos,'status_history':history}
def progress(o):
 expected=int(o['expected'] or 0)
 received=int(o['received'] or 0)
 outbound=int(o['outbound_count'] or 0)
 held=int(o['hold_count'] or 0)
 status=o['goods_status']
 if o.get('return_target',0):
  target=int(o['return_target']);returned=int(o['returned_count']);reout=int(o['reout_count'])
  if reout:return f'重出（{reout}/{target}）'
  return f'退回（{returned}/{target}）'
 if held or o['hold_prealert'] or status=='扣件':return f"{'扣件' if expected>0 and held>=expected else '待扣件'}（{held}/{expected}）"
 if outbound:return f"{'已出库' if expected>0 and outbound>=expected else '待出库'}（{outbound}/{expected}）"
 return f"{'已入库' if expected>0 and received>=expected else '待入库'}（{received}/{expected}）"
def enriched(o):
 d=dict(o);d['progress_status']=progress(d)
 d['hold_label']=bool(d.get('hold_label') or '扣货' in (d.get('remark') or ''))
 d['damage_label']=bool(d.get('damage_label') or '破损' in (d.get('remark') or ''))
 return d
class PieceLookup(BaseModel):
 warehouse_id:str
 barcode:str=Field(min_length=1)
class PieceOperation(BaseModel):
 piece_code:str|None=Field(default=None,max_length=100)
 warehouse_id:str
 barcode:str=Field(min_length=1)
 operation:str
 request_id:str=Field(min_length=1)
 reason:str=''
 damage_choice:str|None=None
 return_all:bool=False
def find_piece_order(c,w,barcode,lock=False):
 # Exact match only. Never silently choose between containers/orders.
 suffix=' FOR UPDATE' if lock and engine.dialect.name=='postgresql' else ''
 rows=c.execute(text('''SELECT DISTINCT o.*,ct.warehouse_id,ct.code container_code,ct.status container_status
 FROM orders o JOIN containers ct ON ct.id=o.container_id
 LEFT JOIN barcodes b ON b.order_id=o.id
 WHERE ct.warehouse_id=:w AND ct.deleted_at IS NULL AND
 (o.system_no=:b OR o.main_barcode=:b OR o.transfer_no=:b OR b.barcode=:b)'''),{'w':w,'b':barcode}).mappings().all()
 if len(rows)!=1:return None,len(rows)
 o=dict(rows[0])
 if lock:
  # Lock the order row and re-read counters inside the transaction.
  c.execute(text('SELECT id FROM orders WHERE id=:id'+suffix),{'id':o['id']}).first()
  refreshed=c.execute(text('SELECT o.*,ct.warehouse_id,ct.code container_code,ct.status container_status FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':o['id']}).mappings().first()
  if not refreshed:return None,0
  o=dict(refreshed)
 return o,1
@app.get('/api/pda/lookup')
def pda_piece_lookup(warehouse_id:str,barcode:str,user=Depends(current)):
 permit(user,'pda')
 barcode=barcode.strip()
 if not barcode:fail('请扫描单号')
 with engine.connect() as c:
  access(c,user,warehouse_id)
  o,n=find_piece_order(c,warehouse_id,barcode)
  if n!=1:fail('未找到订单' if n==0 else '多个货柜或订单匹配，请人工选择',409)
  d=enriched(o);d['auto_hold']=bool(d['hold_prealert'] or d['hold_count'] or d['goods_status']=='扣件')
  return d
@app.post('/api/pda/operation')
def pda_piece_operation(x:PieceOperation,user=Depends(protect)):
 permit(user,'pda')
 barcode=x.barcode.strip();op=x.operation
 if op not in ('入库','出库','扣件','退回','重出','扣货入库','扣货出库'):fail('不支持的操作')
 with engine.begin() as c:
  access(c,user,x.warehouse_id)
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  payload_data=x.model_dump()
  if payload_data.get('piece_code') is None:payload_data.pop('piece_code',None)
  payload=json.dumps(payload_data,sort_keys=True,ensure_ascii=False)
  def replay():
   prior=c.execute(text('SELECT * FROM operation_receipts WHERE request_id=:r'),{'r':x.request_id}).mappings().first()
   if prior:
    if prior['user_id']!=user['id'] or prior['payload']!=payload:fail('请求ID已用于其他操作',409)
    return dict(json.loads(prior['response']),duplicate_request=True)
  cached=replay()
  if cached:return cached
  if c.execute(text('SELECT 1 FROM piece_operations WHERE request_id=:r'),{'r':x.request_id}).first():fail('该请求ID已有历史记录，请查询流水',409)
  o,n=find_piece_order(c,x.warehouse_id,barcode,True)
  if n!=1:fail('未找到唯一订单，禁止自动记件',409)
  cached=replay()
  if cached:return cached
  if o['container_status']=='已关闭':fail('货柜已关闭')
  # Check the actual remarks on the server; a PDA client cannot bypass hold warnings.
  notes=[o['remark'] or '']+[r[0] for r in c.execute(text('SELECT note FROM goods_notes WHERE order_id=:id'),{'id':o['id']})]
  marked_hold=bool(o['hold_label'] or any('扣货' in note for note in notes))
  auto_hold=bool(o['hold_prealert'] or o['hold_count'] or o['goods_status']=='扣件')
  if op in ('扣货入库','扣货出库'):
   if not marked_hold:fail('该订单没有扣货标记，禁止扣货确认操作',409)
   if auto_hold:fail('订单存在旧版扣件流程，请管理员核对后操作',409)
   op='入库' if op=='扣货入库' else '出库'
  elif marked_hold and op not in ('退回','重出'):fail('此票标记为扣货，请使用专用确认入库或确认出库按钮',409)
  if op=='入库' and x.damage_choice not in ('无破损','有破损'):fail('入库前必须选择无破损或有破损',409)
  if auto_hold and op=='入库':op='扣件'
  if not auto_hold and op=='扣件':fail('请先由管理员在货物详情标记预报扣件',409)
  if auto_hold and op in ('出库','退回','重出'):fail('此单处于扣件流程，须管理员解除扣件后操作',409)
  piece_code=(x.piece_code or '').strip()
  if op in ('入库','扣件'):
   if not re.fullmatch(r'(?:WMS|afg|AFG)-[0-9A-Fa-f]{32}',piece_code):fail('请先扫描单号，再扫描本件有效二维码',409)
   if c.execute(text('SELECT 1 FROM piece_labels WHERE code=:code'),{'code':piece_code}).first():fail('此二维码已绑定货物，禁止重复入库',409)
  if op in ('出库','重出'):ensure_damage_repaired(c,o)
  expected=int(o['expected'] or 0);received=int(o['received'] or 0);outbound=int(o['outbound_count'] or 0);held=int(o['hold_count'] or 0)
  if expected<=0:fail('预报件数无效，禁止扫描入库',409)
  reason='';result='成功';old=o['goods_status'];new=old
  if op=='出库' and int(o['return_target'] or 0):fail('已进入退回流程，请使用重出逐件扫码',409)
  if op in ('入库','出库','扣件'):
   counter={'入库':received,'出库':outbound,'扣件':held}[op]
   if op=='出库' and (received<expected or outbound>=received):reason='入库未完成或可出库件数不足'
   elif counter>=expected:reason='超出预报件数'
   if reason:result='异常'
   else:
    if op in ('入库','扣件'):
     try:c.execute(text('INSERT INTO piece_labels(code,order_id,piece_no,operation,user_id) VALUES(:code,:o,:n,:op,:u)'),{'code':piece_code,'o':o['id'],'n':counter+1,'op':op,'u':user['id']})
     except IntegrityError:fail('此二维码已被绑定，禁止重复入库',409)
    field={'入库':'received','出库':'outbound_count','扣件':'hold_count'}[op]
    c.execute(text(f'UPDATE orders SET {field}={field}+1 WHERE id=:id'),{'id':o['id']})
    if op=='入库':
     piece_no=received+1
     existing=c.execute(text('SELECT id FROM order_pieces WHERE order_id=:o AND piece_no=:n'),{'o':o['id'],'n':piece_no}).first()
     if existing:c.execute(text('UPDATE order_pieces SET damage_status=:d WHERE order_id=:o AND piece_no=:n'),{'d':x.damage_choice,'o':o['id'],'n':piece_no})
     else:c.execute(text('INSERT INTO order_pieces(id,order_id,piece_no,damage_status) VALUES(:id,:o,:n,:d)'),{'id':uid(),'o':o['id'],'n':piece_no,'d':x.damage_choice})
     audit(c,user['id'],x.warehouse_id,'PDA逐件破损检查',json.dumps({'order':o['system_no'],'piece_no':piece_no,'damage':x.damage_choice},ensure_ascii=False))
    if counter+1==expected:
     new={'入库':'已入库','出库':'已出库','扣件':'扣件'}[op]
  else:
   if not x.reason.strip():fail('退回、重出必须填写原因')
   target=int(o['return_target'] or 0);returned=int(o['returned_count'] or 0);reout=int(o['reout_count'] or 0);cycle=int(o['return_cycle'] or 0)
   if op=='退回':
    if outbound<=0:fail('必须实际出库后才允许退回',409)
    if target and reout and reout<target:fail('本轮重出未完成，不能开始新退回',409)
    if not target or reout==target:
     target=outbound;returned=0;reout=0;cycle+=1
    if returned>=target:fail('本轮已全部退回，不能重复退回',409)
    if x.damage_choice not in ('无破损','有破损'):fail('退回必须完成一次破损检查',409)
    quantity=target-returned if x.return_all else 1
    returned+=quantity
    c.execute(text('UPDATE orders SET return_target=:t,returned_count=:r,reout_count=0,return_cycle=:cy WHERE id=:id'),{'t':target,'r':returned,'cy':cycle,'id':o['id']})
    c.execute(text('INSERT INTO return_checks(id,order_id,cycle,request_id,quantity,scope,damage_choice,user_id,reason) VALUES(:id,:o,:cy,:req,:q,:scope,:d,:u,:reason)'),{'id':uid(),'o':o['id'],'cy':cycle,'req':x.request_id,'q':quantity,'scope':'all' if x.return_all else 'single','d':x.damage_choice,'u':user['id'],'reason':x.reason.strip()})
    if x.damage_choice=='有破损':c.execute(text('UPDATE orders SET damage_label=1 WHERE id=:id'),{'id':o['id']})
    reason=f'第{cycle}轮退回 {quantity}件；{x.damage_choice}；'+x.reason.strip()
   else:
    if not target or returned!=target:fail('本轮全部退回后才可重出',409)
    if reout>=target:fail('重出件数已达上限',409)
    c.execute(text('UPDATE orders SET reout_count=reout_count+1 WHERE id=:id'),{'id':o['id']})
    reason=f'第{cycle}轮重出 1件；'+x.reason.strip()
   new=op
  if result=='成功' and op=='入库':
   audit(c,user['id'],x.warehouse_id,'PDA入库破损检查',o['system_no']+' '+str(x.damage_choice))
  if new!=old and result=='成功':
   c.execute(text('UPDATE orders SET goods_status=:s WHERE id=:id'),{'s':new,'id':o['id']})
   c.execute(text('INSERT INTO status_history(id,order_id,user_id,old_status,new_status,reason) VALUES(:id,:o,:u,:old,:new,:r)'),{'id':uid(),'o':o['id'],'u':user['id'],'old':old,'new':new,'r':'PDA扫码自动完成' if op in ('入库','出库','扣件') else x.reason.strip()})
  c.execute(text('INSERT INTO piece_operations(id,request_id,order_id,warehouse_id,user_id,barcode,operation,result,reason) VALUES(:id,:r,:o,:w,:u,:b,:op,:result,:reason)'),{'id':uid(),'r':x.request_id,'o':o['id'],'w':x.warehouse_id,'u':user['id'],'b':barcode,'op':op,'result':result,'reason':reason or x.reason.strip()})
  if result=='异常':audit(c,user['id'],x.warehouse_id,'PDA超量/违规扫描',json.dumps({'order':o['system_no'],'operation':op,'reason':reason},ensure_ascii=False))
  updated=dict(c.execute(text('SELECT * FROM orders WHERE id=:id'),{'id':o['id']}).mappings().first())
  response={'result':result,'reason':reason,'operation':op,'order':enriched(updated),'auto_hold':auto_hold}
  c.execute(text('INSERT INTO operation_receipts(request_id,user_id,payload,response) VALUES(:r,:u,:p,:s)'),{'r':x.request_id,'u':user['id'],'p':payload,'s':json.dumps(response,ensure_ascii=False,default=str)})
  return response
class HoldPrealert(BaseModel):
 enabled:bool
 reason:str=''
@app.post('/api/orders/{oid}/hold-prealert')
def set_hold_prealert(oid:str,x:HoldPrealert,user=Depends(protect)):
 permit(user,'repair')
 with engine.begin() as c:
  o=c.execute(text('SELECT o.*,ct.warehouse_id FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('订单不存在',404)
  access(c,user,o['warehouse_id'])
  if not x.enabled and int(o['hold_count'] or 0)>0:fail('已经开始扣件扫描，不能直接解除；请联系管理员处理',409)
  c.execute(text('UPDATE orders SET hold_prealert=:v WHERE id=:id'),{'v':int(x.enabled),'id':oid})
  audit(c,user['id'],o['warehouse_id'],'预报扣件设置',json.dumps({'order':o['system_no'],'enabled':x.enabled,'reason':x.reason},ensure_ascii=False))
 return {'ok':True}

STATUSES=('已入库','扣件','已出库','退回','重出')
class StatusChange(BaseModel):
 status:str
 reason:str=''
@app.post('/api/orders/{oid}/status')
def change_status(oid:str,x:StatusChange,user=Depends(protect)):
 if x.status not in STATUSES:fail('不支持的货物状态')
 if x.status in STATUSES:fail('入库、出库、扣件必须由PDA逐件扫码完成')
 if x.status in ('扣件','退回','重出') and not x.reason.strip():fail('扣件、退回、重出必须填写原因')
 with engine.begin() as c:
  o=c.execute(text('SELECT o.id,o.system_no,o.goods_status,o.received,o.expected,ct.warehouse_id FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL FOR UPDATE'),{'id':oid}).mappings().first()
  if not o:fail('货物不存在',404)
  access(c,user,o['warehouse_id'])
  old=o['goods_status']
  if old==x.status:fail('货物已经是该状态，未重复提交')
  if x.status in ('已出库','重出') and o['received']==0:fail('没有实收记录，禁止出库；请先核对入库扫码')
  if x.status=='重出' and old not in ('已出库','退回','扣件'):fail('重出前必须是已出库、退回或扣件状态')
  c.execute(text('UPDATE orders SET goods_status=:s WHERE id=:id'),{'s':x.status,'id':oid})
  c.execute(text('INSERT INTO status_history(id,order_id,user_id,old_status,new_status,reason) VALUES(:id,:o,:u,:old,:new,:r)'),{'id':uid(),'o':oid,'u':user['id'],'old':old,'new':x.status,'r':x.reason.strip()})
  audit(c,user['id'],o['warehouse_id'],'货物状态修改',json.dumps({'order':o['system_no'],'old':old,'new':x.status,'reason':x.reason.strip()},ensure_ascii=False))
 return {'ok':True,'old_status':old or '待入库','new_status':x.status}
@app.get('/api/search/orders')
def search_orders(warehouse_id:str,q:str,user=Depends(current)):
 permit(user,'read')
 q=q.strip()
 if not q:fail('请输入单号')
 if len(q)>200:fail('单号过长')
 with engine.connect() as c:
  access(c,user,warehouse_id)
  rows=c.execute(text('''SELECT o.id,o.system_no,o.main_barcode,o.transfer_no,o.goods_status,ct.id container_id,ct.code container_code FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE ct.warehouse_id=:w AND ct.deleted_at IS NULL AND (LOWER(o.system_no) LIKE :q ESCAPE '!' OR LOWER(COALESCE(o.main_barcode,'')) LIKE :q ESCAPE '!' OR LOWER(COALESCE(o.transfer_no,'')) LIKE :q ESCAPE '!') ORDER BY CASE WHEN LOWER(o.system_no)=:exact OR LOWER(COALESCE(o.main_barcode,''))=:exact OR LOWER(COALESCE(o.transfer_no,''))=:exact THEN 0 ELSE 1 END,ct.code,o.system_no LIMIT 100'''),{'w':warehouse_id,'q':'%'+q.lower().replace('!','!!').replace('%','!%').replace('_','!_')+'%','exact':q.lower()}).mappings().all()
  return [dict(r) for r in rows]
@app.patch('/api/orders/{oid}')
def edit_order(oid:str,x:OrderEdit,user=Depends(protect)):
 permit(user,'edit')
 with engine.begin() as c:
  o=c.execute(text('SELECT o.*,ct.warehouse_id FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('货物不存在',404)
  access(c,user,o['warehouse_id'])
  changes=x.model_dump(exclude_unset=True)
  if 'expected' in changes and changes['expected'] is None:fail('预报件数不能为空')
  if changes.get('system_no') is not None and c.execute(text('SELECT id FROM orders WHERE container_id=:c AND system_no=:n AND id<>:id'),{'c':o['container_id'],'n':changes['system_no'].strip(),'id':oid}).first():fail('同柜客户单号已存在',409)
  if 'system_no' in changes and changes['system_no'] is not None:changes['system_no']=changes['system_no'].strip()
  if 'expected' in changes and changes['expected']<o['received']:fail('预报件数不能小于已收件数')
  if 'received' in changes and changes['received']!=o['received']:fail('实收件数只能由扫码更新')
  if 'expected' in changes and changes['expected']<=0:fail('预报件数必须大于0')
  if 'expected' in changes and (o['outbound_count'] or o['return_target']) and changes['expected']!=o['expected']:fail('出库后不能修改预报件数')
  if 'mode' in changes and changes['mode'] not in ('待确认','共用','独立'):fail('条码模式无效')
  if 'main_barcode' in changes and changes['main_barcode']!=o['main_barcode'] and o['received']:fail('已有扫码记录，不能直接修改原条码')
  for k,v in changes.items():
   if v is None:
    if k in ('weight','length','width','height','declared_value'):c.execute(text(f'UPDATE orders SET {k}=NULL WHERE id=:id'),{'id':oid})
    continue
   if k=='system_no' and not v.strip():fail('系统单号不能为空')
   c.execute(text(f'UPDATE orders SET {k}=:v WHERE id=:id'),{'v':v,'id':oid})
  if 'remark' in changes:
   if '扣货' in (changes['remark'] or ''):c.execute(text('UPDATE orders SET hold_label=1 WHERE id=:id'),{'id':oid})
   if '破损' in (changes['remark'] or ''):c.execute(text('UPDATE orders SET damage_label=1 WHERE id=:id'),{'id':oid})
  if 'main_barcode' in changes and changes['main_barcode']:
   c.execute(text('DELETE FROM barcodes WHERE order_id=:id'),{'id':oid})
   c.execute(text('INSERT INTO barcodes(id,order_id,barcode) VALUES(:id,:o,:b)'),{'id':uid(),'o':oid,'b':changes['main_barcode']})
  audit(c,user['id'],o['warehouse_id'],'修改货物',json.dumps({'id':oid,'before':{k:o[k] for k in changes},'after':changes},ensure_ascii=False))
 return {'ok':True}
class NewNote(BaseModel):note:str=Field(min_length=1,max_length=5000)
@app.post('/api/orders/{oid}/notes')
def add_note(oid:str,x:NewNote,user=Depends(protect)):
 permit(user,'notes')
 with engine.begin() as c:
  o=c.execute(text('SELECT ct.warehouse_id FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('货物不存在',404)
  access(c,user,o['warehouse_id'])
  c.execute(text('INSERT INTO goods_notes(id,order_id,user_id,note) VALUES(:id,:o,:u,:n)'),{'id':uid(),'o':oid,'u':user['id'],'n':x.note})
  if '扣货' in x.note:c.execute(text('UPDATE orders SET hold_label=1 WHERE id=:id'),{'id':oid})
  if '破损' in x.note:c.execute(text('UPDATE orders SET damage_label=1 WHERE id=:id'),{'id':oid})
  audit(c,user['id'],o['warehouse_id'],'货物备注',oid)
 return {'ok':True}
PHOTO_ROOT=Path(os.getenv('WMS_UPLOAD_DIR',str(ROOT/'uploads')))
@app.post('/api/orders/{oid}/photos')
async def upload_photo(oid:str,file:UploadFile=File(...),user=Depends(protect)):
 permit(user,'notes')
 with engine.connect() as c:
  o=c.execute(text('SELECT ct.warehouse_id FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('货物不存在',404)
  access(c,user,o['warehouse_id'])
 data=await file.read(8_000_001)
 if len(data)>8_000_000:fail('图片不得超过8MB')
 # Verify actual magic bytes; never trust filename or browser-supplied MIME.
 if data.startswith(b'\xff\xd8\xff'):ext,mime='jpg','image/jpeg'
 elif data.startswith(b'\x89PNG\r\n\x1a\n'):ext,mime='png','image/png'
 elif data.startswith(b'RIFF') and data[8:12]==b'WEBP':ext,mime='webp','image/webp'
 else:fail('仅支持 JPG、PNG、WEBP 图片')
 PHOTO_ROOT.mkdir(parents=True,exist_ok=True)
 ident=uid();name=ident+'.'+ext
 (PHOTO_ROOT/name).write_bytes(data)
 with engine.begin() as c:
  c.execute(text('INSERT INTO goods_photos(id,order_id,user_id,filename,mime) VALUES(:id,:o,:u,:f,:m)'),{'id':ident,'o':oid,'u':user['id'],'f':name,'m':mime})
  audit(c,user['id'],o['warehouse_id'],'上传货物照片',oid+' '+ident)
 return {'id':ident}
@app.get('/api/photos/{pid}')
def get_photo(pid:str,user=Depends(current)):
 permit(user,'read')
 with engine.connect() as c:
  p=c.execute(text('SELECT p.*,ct.warehouse_id FROM goods_photos p JOIN orders o ON o.id=p.order_id JOIN containers ct ON ct.id=o.container_id WHERE p.id=:id AND ct.deleted_at IS NULL'),{'id':pid}).mappings().first()
  if not p:fail('照片不存在',404)
  access(c,user,p['warehouse_id'])
  return FileResponse(PHOTO_ROOT/p['filename'],media_type=p['mime'],headers={'Cache-Control':'private, no-store','X-Content-Type-Options':'nosniff'})
class GoodsEvent(BaseModel):kind:str;event_time:str;reason:str=''
@app.post('/api/orders/{oid}/events')
def add_event(oid:str,x:GoodsEvent,user=Depends(protect)):
 permit(user,'edit')
 if x.kind not in ('入库','出库','时间更正'):fail('事件类型无效')
 try:datetime.fromisoformat(x.event_time)
 except ValueError:fail('时间格式错误')
 with engine.begin() as c:
  o=c.execute(text('SELECT ct.warehouse_id FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('货物不存在',404)
  access(c,user,o['warehouse_id'])
  c.execute(text('INSERT INTO goods_events(id,order_id,kind,event_time,user_id,reason) VALUES(:id,:o,:k,:t,:u,:r)'),{'id':uid(),'o':oid,'k':x.kind,'t':x.event_time,'u':user['id'],'r':x.reason})
  audit(c,user['id'],o['warehouse_id'],'货物时间记录',oid+' '+x.kind+' '+x.event_time)
 return {'ok':True}
class NewOrder(BaseModel):container_id:str;system_no:str;expected:int=Field(ge=0);main_barcode:str='';mode:str='待确认';extra_barcodes:list[str]=[]
@app.post('/api/orders')
def new_order(x:NewOrder,user=Depends(protect)):
 permit(user,'forecast')
 if x.mode not in ('待确认','共用','独立'):fail('条码模式无效')
 with engine.begin() as c:
  ct=c.execute(text('SELECT warehouse_id,status FROM containers WHERE id=:id AND deleted_at IS NULL'),{'id':x.container_id}).mappings().first()
  if not ct:fail('货柜不存在',404)
  access(c,user,ct['warehouse_id'])
  if ct['status']=='已关闭':fail('货柜已关闭')
  oid=uid()
  try:c.execute(text('INSERT INTO orders(id,container_id,system_no,main_barcode,expected,mode) VALUES(:id,:c,:s,:b,:e,:m)'),{'id':oid,'c':x.container_id,'s':x.system_no.strip(),'b':x.main_barcode.strip(),'e':x.expected,'m':x.mode})
  except IntegrityError:fail('系统单号重复')
  for b in set([x.main_barcode.strip()]+[v.strip() for v in x.extra_barcodes]):
   if b:c.execute(text('INSERT INTO barcodes(id,order_id,barcode) VALUES(:id,:o,:b)'),{'id':uid(),'o':oid,'b':b})
  audit(c,user['id'],ct['warehouse_id'],'新增预报',x.system_no)
 return {'id':oid}
class Scan(BaseModel):container_id:str;barcode:str=Field(min_length=1);request_id:str=Field(min_length=1)
@app.post('/api/scan')
def scan(x:Scan,user=Depends(protect)):
 with engine.connect() as c:
  ct=c.execute(text('SELECT * FROM containers WHERE id=:id'),{'id':x.container_id}).mappings().first()
  if not ct:fail('货柜不存在',404)
  access(c,user,ct['warehouse_id'])
  prior=c.execute(text('SELECT * FROM scans WHERE request_id=:r'),{'r':x.request_id}).mappings().first()
  if prior and prior['user_id']==user['id'] and prior['container_id']==x.container_id and prior['barcode']==x.barcode.strip():return dict(prior)
  fail('请使用PDA先扫单号，再扫本件二维码后确认入库',409)
class LinkBarcode(BaseModel):order_id:str;mode:str
@app.post('/api/exceptions/{eid}/associate')
def associate(eid:str,x:LinkBarcode,user=Depends(protect)):
 permit(user,'notes')
 if x.mode not in ('共用','独立'):fail('模式不合法')
 with engine.begin() as c:
  ex=c.execute(text('SELECT * FROM exceptions WHERE id=:id'),{'id':eid}).mappings().first()
  if not ex or ex['status']!='待处理':fail('异常不存在或已处理')
  access(c,user,ex['warehouse_id'])
  o=c.execute(text('SELECT * FROM orders WHERE id=:id AND container_id=:cid'),{'id':x.order_id,'cid':ex['container_id']}).mappings().first()
  if not o:fail('订单不属于此柜')
  if o['mode'] not in ('待确认',x.mode):fail('订单模式冲突')
  if o['received']>=o['expected']:fail('预报件数已满')
  c.execute(text('UPDATE orders SET mode=:m WHERE id=:id'),{'m':x.mode,'id':x.order_id})
  if not c.execute(text('SELECT 1 FROM barcodes WHERE order_id=:o AND barcode=:b'),{'o':x.order_id,'b':ex['barcode']}).first():c.execute(text('INSERT INTO barcodes(id,order_id,barcode) VALUES(:id,:o,:b)'),{'id':uid(),'o':x.order_id,'b':ex['barcode']})
  c.execute(text("UPDATE exceptions SET status='已关联待重扫' WHERE id=:id"),{'id':eid})
  audit(c,user['id'],ex['warehouse_id'],'关联异常条码',ex['barcode']+' -> '+x.order_id)
 return {'ok':True,'message':'关联成功，需重新扫描，历史异常不会自动计入库存'}
class Mode(BaseModel):mode:str
@app.post('/api/orders/{oid}/mode')
def mode(oid:str,x:Mode,user=Depends(protect)):
 permit(user,'edit')
 if x.mode not in ('共用','独立','待确认'):fail('模式不合法')
 with engine.begin() as c:
  o=c.execute(text('SELECT o.received,c.warehouse_id FROM orders o JOIN containers c ON c.id=o.container_id WHERE o.id=:id AND c.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('订单不存在')
  access(c,user,o['warehouse_id'])
  if o['received']:fail('已有实收记录，不可直接更改条码模式')
  c.execute(text('UPDATE orders SET mode=:m WHERE id=:id'),{'m':x.mode,'id':oid})
  audit(c,user['id'],o['warehouse_id'],'修改条码模式',oid+' '+x.mode)
 return {'ok':True}
def parse_excel(data,filename):
 if filename.lower().endswith('.xls'):
  book=xlrd.open_workbook(file_contents=data)
  sheets=[(sh.name,[sh.row_values(i) for i in range(sh.nrows)],[(r0+1,r1,c0+1,c1) for r0,r1,c0,c1 in sh.merged_cells]) for sh in book.sheets()]
 else:
  wb=load_workbook(io.BytesIO(data),read_only=False,data_only=True)
  sheets=[(sh.title,list(sh.iter_rows(values_only=True)),[(m.min_row,m.max_row,m.min_col,m.max_col) for m in sh.merged_cells.ranges]) for sh in wb];wb.close()
 def norm(v):return ''.join(str(v if v is not None else '').split()).replace('（','(').replace('）',')').lower()
 def number(v):
  try:
   n=float(v) if str(v).strip() else None
   if n is not None and (not __import__('math').isfinite(n) or n<0):raise ValueError('重量、尺寸、件数和货值不能为负数或无效数值')
   return n
  except (TypeError,ValueError):
   if str(v).strip():raise ValueError('无效数值：'+str(v))
   return None
 result=[]
 aliases={'system_no':['系统单号','客户单号','客户订单号'],'main_barcode':['客户唛头','原运单号','运单号','追踪单号'],'transfer_no':['新标签','转单号','转运单号'],'channel':['渠道'],'recipient':['收件人姓名','收件人'],'company':['公司'],'phone':['联系电话','电话'],'email':['邮箱'],'country':['国家','目的地国家'],'postcode':['邮编'],'city':['城市','收件人城市'],'address':['地址'],'expected':['件数','件数(拖)'],'weight':['实际重量','重量','重量KG','重量(kg)'],'length':['长','长度'],'width':['宽','宽度'],'height':['高','高度'],'declared_value':['货值'],'description':['物品描述'],'remark':['备注']}
 for sheet,rows,merges in sheets:
  header_i=None;columns={}
  for i,row in enumerate(rows[:35]):
   head=[norm(v) for v in row]
   columns={k:next((head.index(norm(a)) for a in names if norm(a) in head),None) for k,names in aliases.items()}
   if columns['system_no'] is not None and columns['expected'] is not None:header_i=i;break
  if header_i is None:continue
  def value(row,k):
   j=columns[k]
   v=row[j] if j is not None and j<len(row) else None
   return str(v).strip() if v is not None else ''
  groups=[];last=None;end_row=None
  for rownum,row in enumerate(rows[header_i+1:],header_i+2):
   sys=value(row,'system_no')
   if norm(sys) in [norm(a) for a in aliases['system_no']]:last=None;continue
   if sys in ('合计','总计','小计'):last=None;continue
   dims={k:number(value(row,k)) for k in ('weight','length','width','height')}
   has_dims=any(v is not None for v in dims.values())
   if not sys:
    if last is not None and has_dims and (end_row is None or rownum<=end_row):
     if len(last['pieces'])>=last['expected']:raise ValueError(f'{sheet}第{rownum}行分件数量超过预报件数')
     last['pieces'].append({**dims,'source_row':rownum})
    elif not has_dims:last=None
    continue
   qty=number(value(row,'expected'))
   if qty is not None and (qty<1 or not qty.is_integer()):raise ValueError(f'{sheet}第{rownum}行件数必须为正整数')
   item={k:value(row,k) for k in aliases if k not in ('expected','weight','length','width','height','declared_value')}
   item.update(expected=int(qty) if qty is not None else 1,declared_value=number(value(row,'declared_value')),pieces=[],**dims)
   if has_dims:item['pieces'].append({**dims,'source_row':rownum})
   groups.append(item);last=item
   end_row=next((r1 for r0,r1,c0,c1 in merges if r0==rownum and c0<=columns['system_no']+1<=c1),None)
   if len(groups)>5000:raise ValueError('单个工作表最多5000单')
  for item in groups:
   if item['expected']>1:
    item['weight']=sum(p['weight'] for p in item['pieces']) if len(item['pieces'])==item['expected'] and all(p['weight'] is not None for p in item['pieces']) else None
    for k in ('length','width','height'):item[k]=None
  if groups:result.append({'sheet':sheet,'rows':groups})
 return {'sheets':result,'total':sum(len(s['rows']) for s in result)}
@app.post('/api/import/preview')
async def preview(file:UploadFile=File(...),user=Depends(protect)):
 permit(user,'forecast');data=await file.read(15_000_001)
 if len(data)>15_000_000:fail('Excel 超过15MB')
 try:parsed=parse_excel(data,file.filename or '')
 except Exception as e:fail('Excel 解析失败：'+str(e)[:140])
 if not parsed['total']:fail('没有识别到客户单号/系统单号及件数')
 with engine.begin() as c:parsed['source_id']=store_original_forecast(c,data,file.filename or '',user)
 return parsed
def import_rows_in_transaction(c,x,user):
 ct=c.execute(text('SELECT * FROM containers WHERE id=:id AND deleted_at IS NULL'),{'id':x.container_id}).mappings().first()
 if not ct:fail('柜号不存在')
 access(c,user,ct['warehouse_id'])
 if ct['status']=='已关闭':fail('货柜已关闭')
 if not ct['customer_id'] or x.customer_id!=ct['customer_id']:fail('客户标签与货柜不一致')
 imported=0;skipped=0;updated=0
 for r in x.rows:
  sys=str(r.get('system_no','')).strip()
  try:qty=int(r.get('expected',0))
  except (ValueError,TypeError):skipped+=1;continue
  if not sys or qty<=0:skipped+=1;continue
  existing=c.execute(text('SELECT * FROM orders WHERE container_id=:c AND system_no=:s'+(' FOR UPDATE' if engine.dialect.name=='postgresql' else '')),{'c':x.container_id,'s':sys}).mappings().first()
  if existing and qty!=existing['expected']:fail('已有订单 '+sys+' 件数与Excel不同，请先核对，未覆盖历史')
  if len(r.get('pieces') or [])>qty:fail('分件数量超过预报件数：'+sys)
  oid=existing['id'] if existing else uid();fields=['id','container_id','system_no','main_barcode','transfer_no','expected','mode','channel','country','city','postcode','weight','length','width','height','recipient','company','phone','email','address','declared_value','description','remark'];params={'id':oid,'container_id':x.container_id,'system_no':sys,'expected':qty,'mode':'待确认'}
  for k in fields:params.setdefault(k,r.get(k) if r.get(k) is not None and r.get(k)!='' else None)
  if existing:
   # Only enrich supplied reference data; counters, damage and history are untouched.
   changes={k:params[k] for k in fields if k not in ('id','container_id','system_no','expected','mode') and params[k] is not None}
   if existing['hold_label'] and 'remark' in changes and '扣货' not in changes['remark']:changes.pop('remark')
   if changes:c.execute(text('UPDATE orders SET '+','.join(k+'=:'+k for k in changes)+' WHERE id=:id'),dict(changes,id=oid))
   audit(c,user['id'],ct['warehouse_id'],'Excel更新已有订单资料',json.dumps({'order':sys,'before':{k:existing[k] for k in changes},'after':changes},ensure_ascii=False,default=str))
   updated+=1
  else:c.execute(text('INSERT INTO orders('+','.join(fields)+') VALUES('+','.join(':'+k for k in fields)+')'),params)
  bc=str(r.get('main_barcode') or '').strip()
  if bc and not c.execute(text('SELECT id FROM barcodes WHERE order_id=:o AND barcode=:b'),{'o':oid,'b':bc}).first():c.execute(text('INSERT INTO barcodes(id,order_id,barcode) VALUES(:id,:o,:b)'),{'id':uid(),'o':oid,'b':bc})
  for j,piece in enumerate((r.get('pieces') or [])[:qty],1):
   prior=c.execute(text('SELECT id FROM order_pieces WHERE order_id=:o AND piece_no=:n'),{'o':oid,'n':j}).first()
   if prior:
    dims={k:piece[k] for k in ('weight','length','width','height','source_row') if piece.get(k) is not None}
    if dims:c.execute(text('UPDATE order_pieces SET '+','.join(k+'=:'+k for k in dims)+' WHERE order_id=:o AND piece_no=:n'),dict(dims,o=oid,n=j))
   else:
    c.execute(text('INSERT INTO order_pieces(id,order_id,piece_no,weight,length,width,height,source_row) VALUES(:id,:o,:n,:w,:l,:b,:h,:r)'),{'id':uid(),'o':oid,'n':j,'w':piece.get('weight'),'l':piece.get('length'),'b':piece.get('width'),'h':piece.get('height'),'r':piece.get('source_row')})
  if not existing:imported+=1
 audit(c,user['id'],ct['warehouse_id'],'导入Excel',f'新增{imported} 更新{updated} 跳过{skipped}')
 return {'imported':imported,'updated':updated,'skipped':skipped}

class ImportCommit(BaseModel):container_id:str;rows:list[dict];customer_id:str=''
@app.post('/api/import/commit')
def commit(x:ImportCommit,user=Depends(protect)):
 permit(user,'forecast')
 if len(x.rows)>5000:fail('单次导入最多5000单')
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  return import_rows_in_transaction(c,x,user)
class PieceEdit(BaseModel):weight:float|None=None;length:float|None=None;width:float|None=None;height:float|None=None;note:str|None=None
class PieceDamageEdit(BaseModel):damage_status:str
@app.patch('/api/orders/{oid}/pieces/{piece_no}')
def edit_piece(oid:str,piece_no:int,x:PieceEdit,user=Depends(protect)):
 permit(user,'edit')
 with engine.begin() as c:
  o=c.execute(text('SELECT o.*,ct.warehouse_id FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('货物不存在',404)
  access(c,user,o['warehouse_id'])
  if piece_no<1 or piece_no>o['expected']:fail('件号超出预报件数')
  before=c.execute(text('SELECT * FROM order_pieces WHERE order_id=:o AND piece_no=:n'),{'o':oid,'n':piece_no}).mappings().first()
  if not before:c.execute(text('INSERT INTO order_pieces(id,order_id,piece_no) VALUES(:id,:o,:n)'),{'id':uid(),'o':oid,'n':piece_no})
  changes=x.model_dump(exclude_unset=True)
  for k,v in changes.items():c.execute(text(f'UPDATE order_pieces SET {k}=:v WHERE order_id=:o AND piece_no=:n'),{'v':v,'o':oid,'n':piece_no})
  audit(c,user['id'],o['warehouse_id'],'修改单件货物',json.dumps({'order':oid,'piece':piece_no,'before':dict(before) if before else None,'after':changes},ensure_ascii=False,default=str))
 return {'ok':True}
@app.patch('/api/orders/{oid}/pieces/{piece_no}/damage')
def edit_piece_damage(oid:str,piece_no:int,x:PieceDamageEdit,user=Depends(protect)):
 permit(user,'repair')
 if x.damage_status!='已修复':fail('仅允许确认已修复')
 with engine.begin() as c:
  o=c.execute(text('SELECT o.*,ct.warehouse_id FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('货物不存在',404)
  access(c,user,o['warehouse_id'])
  if piece_no<1 or piece_no>int(o['received'] or 0):fail('只能修改已入库货物的破损状态')
  before=c.execute(text('SELECT damage_status FROM order_pieces WHERE order_id=:o AND piece_no=:n'),{'o':oid,'n':piece_no}).first()
  if not before or before[0]!='有破损':fail('只有有破损的货物才能确认已修复',409)
  c.execute(text('UPDATE order_pieces SET damage_status=:v WHERE order_id=:o AND piece_no=:n'),{'v':'已修复','o':oid,'n':piece_no})
  audit(c,user['id'],o['warehouse_id'],'确认单件破损已修复',json.dumps({'order':o['system_no'],'piece_no':piece_no,'old':before[0] if before else None,'new':x.damage_status},ensure_ascii=False))
 return {'ok':True}
@app.post('/api/orders/{oid}/return-checks/{check_id}/repair')
def repair_return_damage(oid:str,check_id:str,user=Depends(protect)):
 permit(user,'repair')
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  suffix=' FOR UPDATE' if engine.dialect.name=='postgresql' else ''
  row=c.execute(text('SELECT id FROM orders WHERE id=:id'+suffix),{'id':oid}).first()
  if not row:fail('货物不存在',404)
  o=c.execute(text('SELECT o.*,ct.warehouse_id FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
  if not o:fail('货物不存在',404)
  access(c,user,o['warehouse_id'])
  check=c.execute(text('SELECT * FROM return_checks WHERE id=:id AND order_id=:o'),{'id':check_id,'o':oid}).mappings().first()
  if not check:fail('退回检查记录不存在',404)
  if check['damage_choice']!='有破损':fail('该记录没有破损，不能确认修复',409)
  if check['repaired_at']:return {'ok':True,'already_repaired':True}
  stamp=datetime.now(timezone.utc).isoformat()
  c.execute(text('UPDATE return_checks SET repaired_at=:t,repaired_by=:u WHERE id=:id'),{'t':stamp,'u':user['id'],'id':check_id})
  audit(c,user['id'],o['warehouse_id'],'确认退回破损已修复',json.dumps({'order':o['system_no'],'check_id':check_id,'cycle':check['cycle'],'quantity':check['quantity'],'repaired_at':stamp},ensure_ascii=False))
 return {'ok':True}
def warehouse_capacity(c,warehouse_id):
 total_received=0;current_outbound=0;in_stock=0;pending_returns=0
 rows=c.execute(text('SELECT o.received,o.hold_count,o.outbound_count,o.returned_count,o.reout_count FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE ct.warehouse_id=:w AND ct.deleted_at IS NULL'),{'w':warehouse_id}).mappings()
 for row in rows:
  received=max(0,int(row['received'] or 0),int(row['hold_count'] or 0))
  pending=max(0,int(row['returned_count'] or 0)-int(row['reout_count'] or 0))
  outside=max(0,int(row['outbound_count'] or 0)-pending)
  total_received+=received;current_outbound+=outside;in_stock+=max(0,received-outside);pending_returns+=pending
 threshold=200
 return {'received':total_received,'outbound':current_outbound,'in_stock':in_stock,'pending_returns':pending_returns,'threshold':threshold,'over':max(0,in_stock-threshold),'is_full':in_stock>threshold,'usage_percent':round(in_stock/threshold*100,1)}

@app.get('/api/dashboard')
def dashboard(warehouse_id:str,user=Depends(current)):
 permit(user,'read')
 with engine.connect() as c:
  access(c,user,warehouse_id)
  summary=dict(c.execute(text('''SELECT COUNT(DISTINCT ct.id) containers,COALESCE(SUM(o.expected),0) expected,COALESCE(SUM(o.received),0) received FROM containers ct LEFT JOIN orders o ON o.container_id=ct.id WHERE ct.warehouse_id=:w AND ct.deleted_at IS NULL'''),{'w':warehouse_id}).mappings().first())
  summary['capacity']=warehouse_capacity(c,warehouse_id)
  summary['pending_exceptions']=c.execute(text("SELECT COUNT(*) FROM exceptions WHERE warehouse_id=:w AND status='待处理'"),{'w':warehouse_id}).scalar()
  summary['recent']=[dict(r) for r in c.execute(text('SELECT a.action,a.detail,a.created_at,u.username FROM audit_logs a JOIN users u ON u.id=a.user_id WHERE a.warehouse_id=:w ORDER BY a.created_at DESC LIMIT 20'),{'w':warehouse_id}).mappings()]
  for entry in summary['recent']:entry['detail_display']=readable_audit(c,entry['action'],entry['detail'],warehouse_id)
  return summary
@app.get('/health')
def health():return {'ok':True}
@app.get('/')
def index():return FileResponse(ROOT/'static'/'index.html')

# Inserted into app.py by build_client_release.py. This file is not a separate runtime dependency.
def client_company(c,user):
 allow(user,['客户'])
 company=c.execute(text('SELECT id,name FROM customers WHERE id=:id'),{'id':user.get('customer_id')}).mappings().first()
 if not company:fail('客户账号尚未绑定公司，请联系管理员',403)
 return dict(company)

def client_draft(c,user,did,lock=False):
 company=client_company(c,user)
 suffix=' FOR UPDATE' if lock and engine.dialect.name=='postgresql' else ''
 d=c.execute(text('SELECT * FROM client_forecasts WHERE id=:id AND customer_id=:cu'+suffix),{'id':did,'cu':company['id']}).mappings().first()
 if not d:fail('预报不存在',404)
 return dict(d)

def forecast_json(d):
 out=dict(d);out['rows']=json.loads(out.pop('rows_json'))
 return out

def validate_client_rows(rows):
 if not rows or len(rows)>5000:fail('预报必须包含1至5000票货物')
 clean=[];seen=set()
 strings=('system_no','main_barcode','transfer_no','channel','recipient','company','phone','email','country','postcode','city','address','description','remark')
 def positive_number(v):
  if v is None or v=='':return None
  try:n=float(v)
  except (ValueError,TypeError):fail('重量、尺寸或货值无效')
  if not __import__('math').isfinite(n) or n<0:fail('重量、尺寸或货值必须为有效非负数')
  return n
 for row in rows:
  if not isinstance(row,dict):fail('订单资料格式无效')
  r={k:str(row.get(k) or '').strip() for k in strings}
  if not r['system_no'] or len(r['system_no'])>150:fail('客户单号不能为空且最多150字')
  if r['system_no'] in seen:fail('预报包含重复客户单号：'+r['system_no'])
  seen.add(r['system_no'])
  if any(len(r[k])>200 for k in ('main_barcode','transfer_no')) or any(len(v)>5000 for v in r.values()):fail('预报文字字段过长')
  try:q=float(row.get('expected',0))
  except (ValueError,TypeError):fail('件数无效')
  if not q.is_integer() or not 1<=q<=5000:fail('每票件数必须为1至5000的整数')
  r['expected']=int(q);pieces=row.get('pieces') or []
  if not isinstance(pieces,list) or len(pieces)>r['expected']:fail('分件明细数量超过预报件数')
  r['pieces']=[]
  for p in pieces:
   if not isinstance(p,dict):fail('分件资料格式无效')
   item={k:positive_number(p.get(k)) for k in ('weight','length','width','height')}
   source=p.get('source_row')
   if source is not None:
    try:item['source_row']=int(source)
    except (ValueError,TypeError):fail('分件来源行号无效')
   r['pieces'].append(item)
  for k in ('weight','length','width','height','declared_value'):r[k]=positive_number(row.get(k))
  if r['expected']>1:
   r['weight']=sum(p['weight'] for p in r['pieces']) if len(r['pieces'])==r['expected'] and all(p['weight'] is not None for p in r['pieces']) else None
   for k in ('length','width','height'):r[k]=None
  clean.append(r)
 return clean

class ClientAccount(BaseModel):
 username:str=Field(min_length=2,max_length=100)
 password:str=Field(min_length=8,max_length=200)
 customer_id:str
 warehouse_id:str

@app.post('/api/client-accounts')
def create_client_account(x:ClientAccount,user=Depends(protect)):
 allow(user,['老板'])
 if len(x.username.strip())<2:fail('账号去除空格后至少2个字符')
 with engine.begin() as c:
  access(c,user,x.warehouse_id)
  if not c.execute(text('SELECT id FROM customers WHERE id=:id AND deleted_at IS NULL'),{'id':x.customer_id}).first():fail('客户公司不存在')
  ident=uid()
  try:c.execute(text("INSERT INTO users(id,username,password_hash,role,customer_id,client_warehouse_id) VALUES(:id,:u,:p,'客户',:cu,:w)"),{'id':ident,'u':x.username.strip(),'p':pw_hash(x.password),'cu':x.customer_id,'w':x.warehouse_id})
  except IntegrityError:fail('账号已存在',409)
  audit(c,user['id'],x.warehouse_id,'创建客户账号',json.dumps({'username':x.username,'customer_id':x.customer_id},ensure_ascii=False))
 return {'id':ident}

@app.get('/api/client/profile')
def client_profile(user=Depends(current)):
 with engine.connect() as c:
  company=client_company(c,user)
  warehouse=c.execute(text('SELECT id,name FROM warehouses WHERE id=:id'),{'id':user.get('client_warehouse_id')}).mappings().first()
  if not warehouse:fail('客户账号尚未绑定接收仓库，请联系管理员',403)
  return {'username':user['username'],'company':company,'warehouse':dict(warehouse)}

@app.get('/api/client/forecasts')
def client_forecasts(user=Depends(current)):
 with engine.connect() as c:
  company=client_company(c,user)
  return [forecast_json(d) for d in c.execute(text('SELECT * FROM client_forecasts WHERE customer_id=:cu ORDER BY updated_at DESC,id'),{'cu':company['id']}).mappings()]

@app.post('/api/client/forecasts')
def new_client_forecast(user=Depends(protect)):
 with engine.begin() as c:
  company=client_company(c,user);wid=user.get('client_warehouse_id')
  if not c.execute(text('SELECT id FROM warehouses WHERE id=:id'),{'id':wid}).first():fail('接收仓库不存在')
  did=uid();stamp=datetime.now(timezone.utc).isoformat()
  c.execute(text("INSERT INTO client_forecasts(id,customer_id,warehouse_id,code,status,filename,rows_json,revision,created_by,created_at,updated_at) VALUES(:id,:cu,:w,'','draft','','[]',1,:u,:t,:t)"),{'id':did,'cu':company['id'],'w':wid,'u':user['id'],'t':stamp})
  return forecast_json(c.execute(text('SELECT * FROM client_forecasts WHERE id=:id'),{'id':did}).mappings().first())

class ClientDraftEdit(BaseModel):
 source_id:str|None=Field(default=None,max_length=36)
 code:str=Field(default='',max_length=120)
 filename:str=Field(default='',max_length=255)
 rows:list[dict]=Field(default_factory=list)
 revision:int=Field(ge=1)

@app.patch('/api/client/forecasts/{did}')
def save_client_forecast(did:str,x:ClientDraftEdit,user=Depends(protect)):
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  d=client_draft(c,user,did,True)
  if d['status']!='draft':fail('正式预报已锁定，不可修改',409)
  if d['revision']!=x.revision:fail('草稿已被本公司其他账号更新，请刷新后重新打开',409)
  rows=validate_client_rows(x.rows) if x.rows else []
  if x.source_id:
   source=verify_original_source(c,x.source_id,user,d['customer_id'])
   if source['container_id']:fail('该原表已经用于正式货柜',409)
   c.execute(text('UPDATE client_forecasts SET source_id=:s WHERE id=:id'),{'s':x.source_id,'id':did})
  if len(json.dumps(rows,ensure_ascii=False))>15_000_000:fail('预报资料过大')
  c.execute(text('UPDATE client_forecasts SET code=:code,filename=:f,rows_json=:r,revision=revision+1,updated_at=:t WHERE id=:id'),{'code':x.code.strip().upper(),'f':x.filename,'r':json.dumps(rows,ensure_ascii=False),'t':datetime.now(timezone.utc).isoformat(),'id':did})
  return forecast_json(c.execute(text('SELECT * FROM client_forecasts WHERE id=:id'),{'id':did}).mappings().first())

@app.delete('/api/client/forecasts/{did}')
def delete_client_forecast(did:str,user=Depends(protect)):
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  d=client_draft(c,user,did,True)
  if d['status']!='draft':fail('正式预报已锁定，不可删除',409)
  c.execute(text('DELETE FROM client_forecasts WHERE id=:id'),{'id':did})
 return {'ok':True}

@app.post('/api/client/parse')
async def client_parse(file:UploadFile=File(...),user=Depends(protect)):
 with engine.connect() as c:client_company(c,user)
 if not (file.filename or '').lower().endswith(('.xlsx','.xls')):fail('请上传 xlsx 或 xls 文件')
 data=await file.read(15_000_001)
 if len(data)>15_000_000:fail('Excel 超过15MB')
 try:parsed=parse_excel(data,file.filename or '')
 except Exception as e:fail('Excel 解析失败：'+str(e)[:140])
 rows=[r for sh in parsed['sheets'] for r in sh['rows']]
 clean=validate_client_rows(rows)
 with engine.begin() as c:source_id=store_original_forecast(c,data,file.filename or '',user,user['customer_id'])
 return {'rows':clean,'source_id':source_id}

class ClientSubmit(BaseModel):revision:int=Field(ge=1)

@app.post('/api/client/forecasts/{did}/submit')
def submit_client_forecast(did:str,x:ClientSubmit,user=Depends(protect)):
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  d=client_draft(c,user,did,True)
  if d['status']=='void':fail('该货柜已作废，不能再次提交',409)
  # Retries return the original outcome instead of creating duplicate stock.
  if d['status']=='submitted':return {'ok':True,'duplicate_request':True,'forecast':forecast_json(d)}
  if d['revision']!=x.revision:fail('草稿已更新，请刷新核对后提交',409)
  if not d['code'].strip():fail('请先填写柜号并保存草稿')
  rows=validate_client_rows(json.loads(d['rows_json']))
  suffix=' FOR UPDATE' if engine.dialect.name=='postgresql' else ''
  if not c.execute(text('SELECT id FROM warehouses WHERE id=:id'+suffix),{'id':d['warehouse_id']}).first():fail('接收仓库不存在')
  if c.execute(text('SELECT id FROM containers WHERE warehouse_id=:w AND UPPER(code)=:code'),{'w':d['warehouse_id'],'code':d['code'].upper()}).first():fail('该柜号已存在，请联系公司端核对，不能覆盖已有正式货柜',409)
  cid=uid();stamp=datetime.now(timezone.utc).isoformat()
  c.execute(text("INSERT INTO containers(id,warehouse_id,code,customer_id,status) VALUES(:id,:w,:code,:cu,'预报')"),{'id':cid,'w':d['warehouse_id'],'code':d['code'],'cu':d['customer_id']})
  fields=('system_no','main_barcode','transfer_no','channel','recipient','company','phone','email','country','postcode','city','address','description','remark','expected','weight','length','width','height','declared_value')
  for r in rows:
   oid=uid();params={k:r.get(k) for k in fields};params.update(id=oid,container_id=cid)
   cols=('id','container_id')+fields
   c.execute(text('INSERT INTO orders('+','.join(cols)+",mode) VALUES("+','.join(':'+k for k in cols)+",'共用')"),params)
   for bc in set(filter(None,(r['main_barcode'],r['transfer_no']))):c.execute(text('INSERT INTO barcodes(id,order_id,barcode) VALUES(:id,:o,:b)'),{'id':uid(),'o':oid,'b':bc})
   for n,p in enumerate(r['pieces'],1):c.execute(text('INSERT INTO order_pieces(id,order_id,piece_no,weight,length,width,height,source_row) VALUES(:id,:o,:n,:w,:l,:b,:h,:r)'),{'id':uid(),'o':oid,'n':n,'w':p['weight'],'l':p['length'],'b':p['width'],'h':p['height'],'r':p.get('source_row')})
   if '扣货' in r['remark']:c.execute(text('UPDATE orders SET hold_label=1 WHERE id=:id'),{'id':oid})
   if '破损' in r['remark']:c.execute(text('UPDATE orders SET damage_label=1 WHERE id=:id'),{'id':oid})
  bind_original_forecast(c,d.get('source_id'),cid,user,d['customer_id'])
  c.execute(text("UPDATE client_forecasts SET status='submitted',container_id=:c,submitted_by=:u,submitted_at=:t,updated_at=:t,revision=revision+1 WHERE id=:id"),{'c':cid,'u':user['id'],'t':stamp,'id':did})
  audit(c,user['id'],d['warehouse_id'],'客户提交正式预报',json.dumps({'customer_id':d['customer_id'],'code':d['code'],'orders':len(rows),'quantity':sum(r['expected'] for r in rows)},ensure_ascii=False))
  return {'ok':True,'forecast':forecast_json(c.execute(text('SELECT * FROM client_forecasts WHERE id=:id'),{'id':did}).mappings().first())}

def client_progress_text(expected,received,outbound,pending):
 expected=max(0,int(expected or 0));received=max(0,int(received or 0));outbound=max(0,int(outbound or 0));pending=max(0,int(pending or 0))
 if pending:return f'退回待处理：{pending}件'
 if expected and outbound>=expected:return '全部出库'
 if outbound:return f'出库中（{outbound}/{expected}）'
 if expected and received>=expected:return '全部入库'
 if received:return f'入库中（{received}/{expected}）'
 return '已预报'

def client_public_order(o):
 public=enriched(o);public['remark']=''
 pending=max(0,int(public.get('returned_count') or 0)-int(public.get('reout_count') or 0))
 outbound=max(0,int(public.get('outbound_count') or 0)-pending)
 public['pending_returns']=pending;public['current_outbound']=outbound
 public['progress_status']=client_progress_text(public['expected'],public['received'],outbound,pending)
 if not pending and (public.get('hold_count') or public.get('hold_prealert') or public.get('goods_status')=='扣件'):public['progress_status']=progress(public)
 return public

def client_order(c,user,oid):
 company=client_company(c,user)
 o=c.execute(text('SELECT o.*,ct.code container_code FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.customer_id=:cu AND ct.deleted_at IS NULL'),{'id':oid,'cu':company['id']}).mappings().first()
 if not o:fail('货物不存在',404)
 # Warehouse internal notes, employee identities, scans and audit logs are not exposed.
 return client_public_order(o)

@app.get('/api/client/containers')
def client_containers(user=Depends(current)):
 with engine.connect() as c:
  company=client_company(c,user)
  rows=c.execute(text("SELECT ct.id,ct.code,ct.status warehouse_status,ct.created_at,COALESCE(SUM(o.expected),0) expected,COALESCE(SUM(o.received),0) received,COALESCE(SUM(o.outbound_count),0) outbound_total,COALESCE(SUM(CASE WHEN COALESCE(o.returned_count,0)>COALESCE(o.reout_count,0) THEN COALESCE(o.returned_count,0)-COALESCE(o.reout_count,0) ELSE 0 END),0) pending_returns FROM containers ct LEFT JOIN orders o ON o.container_id=ct.id WHERE ct.customer_id=:cu AND ct.deleted_at IS NULL GROUP BY ct.id,ct.code,ct.status,ct.created_at ORDER BY ct.created_at DESC"),{'cu':company['id']}).mappings()
  result=[]
  for row in rows:
   d=dict(row);d['outbound']=max(0,int(d['outbound_total'])-int(d['pending_returns']))
   d['status']=client_progress_text(d['expected'],d['received'],d['outbound'],d['pending_returns']);result.append(d)
  return result

@app.get('/api/client/orders')
def client_orders(q:str='',container_id:str='',user=Depends(current)):
 with engine.connect() as c:
  company=client_company(c,user)
  if container_id and not c.execute(text('SELECT id FROM containers WHERE id=:id AND customer_id=:cu AND deleted_at IS NULL'),{'id':container_id,'cu':company['id']}).first():fail('货柜不存在',404)
  rows=c.execute(text("SELECT o.*,ct.code container_code FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE ct.customer_id=:cu AND ct.deleted_at IS NULL AND (:cid='' OR ct.id=:cid) AND (:q='' OR LOWER(o.system_no) LIKE :like OR LOWER(COALESCE(o.main_barcode,'')) LIKE :like OR LOWER(COALESCE(o.transfer_no,'')) LIKE :like OR LOWER(ct.code) LIKE :like) ORDER BY ct.created_at DESC,o.system_no LIMIT 5000"),{'cu':company['id'],'cid':container_id,'q':q.strip(),'like':'%'+q.strip().lower()+'%'}).mappings()
  result=[]
  for r in rows:
   result.append(client_public_order(r))
  return result

@app.get('/api/client/orders/{oid}')
def client_order_detail(oid:str,user=Depends(current)):
 with engine.connect() as c:
  o=client_order(c,user,oid)
  pieces=[dict(p) for p in c.execute(text('SELECT piece_no,weight,length,width,height,damage_status FROM order_pieces WHERE order_id=:id ORDER BY piece_no'),{'id':oid}).mappings()]
  checks=[dict(p) for p in c.execute(text('SELECT cycle,quantity,damage_choice,repaired_at,created_at FROM return_checks WHERE order_id=:id ORDER BY created_at DESC'),{'id':oid}).mappings()]
  photos=[dict(p) for p in c.execute(text('SELECT id,created_at FROM goods_photos WHERE order_id=:id ORDER BY created_at DESC'),{'id':oid}).mappings()]
  events=[dict(p) for p in c.execute(text("SELECT kind,event_time FROM goods_events WHERE order_id=:id AND kind IN ('入库','出库') ORDER BY event_time DESC"),{'id':oid}).mappings()]
  return {'order':o,'pieces':pieces,'return_checks':checks,'photos':photos,'events':events}

@app.get('/api/client/photos/{pid}')
def client_photo(pid:str,user=Depends(current)):
 with engine.connect() as c:
  company=client_company(c,user)
  p=c.execute(text('SELECT p.filename,p.mime FROM goods_photos p JOIN orders o ON o.id=p.order_id JOIN containers ct ON ct.id=o.container_id WHERE p.id=:id AND ct.customer_id=:cu AND ct.deleted_at IS NULL'),{'id':pid,'cu':company['id']}).mappings().first()
  if not p:fail('照片不存在',404)
  path=(PHOTO_ROOT/p['filename']).resolve()
  if PHOTO_ROOT.resolve() not in path.parents:fail('照片不存在',404)
  return FileResponse(path,media_type=p['mime'],headers={'Cache-Control':'private, no-store','X-Content-Type-Options':'nosniff'})

@app.get('/client')
def client_page():return FileResponse(ROOT/'static'/'customer.html',headers={'Cache-Control':'no-store'})

class CreateContainerImport(BaseModel):
 source_id:str=Field(default='',max_length=36)
 warehouse_id:str
 customer_id:str
 code:str=Field(min_length=1,max_length=120)
 rows:list[dict]
 request_id:str=Field(min_length=1,max_length=100)

@app.post('/api/import/create-container')
def create_container_import(x:CreateContainerImport,user=Depends(protect)):
 permit(user,'forecast')
 code=x.code.strip().upper()
 if not code:fail('柜号不能为空')
 rows=validate_client_rows(x.rows)
 fingerprint_data=x.model_dump()
 if not fingerprint_data.get('source_id'):fingerprint_data.pop('source_id',None)
 fingerprint=hashlib.sha256(json.dumps(fingerprint_data,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  access(c,user,x.warehouse_id)
  suffix=' FOR UPDATE' if engine.dialect.name=='postgresql' else ''
  if not c.execute(text('SELECT id FROM warehouses WHERE id=:id'+suffix),{'id':x.warehouse_id}).first():fail('仓库不存在')
  prior=c.execute(text('SELECT * FROM container_import_receipts WHERE request_id=:id'),{'id':x.request_id}).mappings().first()
  if prior:
   if prior['user_id']!=user['id'] or prior['payload_hash']!=fingerprint:fail('请求ID已用于其他导入',409)
   return dict(json.loads(prior['response']),duplicate_request=True)
  if not c.execute(text('SELECT id FROM customers WHERE id=:id AND deleted_at IS NULL'),{'id':x.customer_id}).first():fail('请选择有效客户')
  if c.execute(text('SELECT id FROM containers WHERE warehouse_id=:w AND UPPER(code)=:code'),{'w':x.warehouse_id,'code':code}).first():fail('柜号已存在，请从货柜列表查看现有货柜',409)
  cid=uid()
  c.execute(text("INSERT INTO containers(id,warehouse_id,code,customer_id,status) VALUES(:id,:w,:code,:cu,'预报')"),{'id':cid,'w':x.warehouse_id,'code':code,'cu':x.customer_id})
  result=import_rows_in_transaction(c,ImportCommit(container_id=cid,customer_id=x.customer_id,rows=rows),user)
  bind_original_forecast(c,x.source_id,cid,user)
  result['container_id']=cid;result['code']=code
  audit(c,user['id'],x.warehouse_id,'创建货柜并导入预报',json.dumps({'code':code,'customer_id':x.customer_id,'orders':len(rows)},ensure_ascii=False))
  c.execute(text('INSERT INTO container_import_receipts(request_id,user_id,payload_hash,response) VALUES(:r,:u,:p,:s)'),{'r':x.request_id,'u':user['id'],'p':fingerprint,'s':json.dumps(result,ensure_ascii=False)})
  return result

class WarehouseMembershipEdit(BaseModel):enabled:bool

@app.patch('/api/warehouses/{warehouse_id}')
def rename_warehouse(warehouse_id:str,x:NewWarehouse,user=Depends(protect)):
 permit(user,'warehouses')
 name=x.name.strip()
 if not name:fail('仓库名称不能为空')
 with engine.begin() as c:
  access(c,user,warehouse_id)
  old=c.execute(text('SELECT name FROM warehouses WHERE id=:id'),{'id':warehouse_id}).scalar()
  if old is None:fail('仓库不存在',404)
  if c.execute(text('SELECT id FROM warehouses WHERE LOWER(name)=LOWER(:n) AND id<>:id'),{'n':name,'id':warehouse_id}).first():fail('仓库名称已存在',409)
  c.execute(text('UPDATE warehouses SET name=:n WHERE id=:id'),{'n':name,'id':warehouse_id})
  audit(c,user['id'],warehouse_id,'修改仓库名称',json.dumps({'old':old,'new':name},ensure_ascii=False))
 return {'ok':True}

@app.get('/api/warehouses/{warehouse_id}/members')
def warehouse_members(warehouse_id:str,user=Depends(current)):
 permit(user,'warehouses')
 with engine.connect() as c:
  access(c,user,warehouse_id)
  return [dict(r) for r in c.execute(text("SELECT u.id,u.username,u.role,u.active,CASE WHEN m.user_id IS NULL THEN 0 ELSE 1 END assigned FROM users u LEFT JOIN memberships m ON m.user_id=u.id AND m.warehouse_id=:w WHERE u.role IN ('老板','主管','客服','仓库管理员','PDA操作员','PDA员工','只读人员') ORDER BY u.username"),{'w':warehouse_id}).mappings()]

@app.put('/api/warehouses/{warehouse_id}/members/{target}')
def set_warehouse_member(warehouse_id:str,target:str,x:WarehouseMembershipEdit,user=Depends(protect)):
 permit(user,'warehouses')
 if target==user['id'] and not x.enabled:fail('不能撤销自己对当前仓库的权限')
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  access(c,user,warehouse_id)
  suffix=' FOR UPDATE' if engine.dialect.name=='postgresql' else ''
  if not c.execute(text('SELECT id FROM warehouses WHERE id=:id'+suffix),{'id':warehouse_id}).first():fail('仓库不存在',404)
  employee=c.execute(text('SELECT id,role FROM users WHERE id=:id'),{'id':target}).mappings().first()
  if not employee or employee['role'] not in ('老板','主管','客服','仓库管理员','PDA操作员','PDA员工','只读人员'):fail('请选择公司员工，客户账号不分配公司端仓库权限')
  if user['role']!='老板' and employee['role']=='老板':fail('主管不能修改老板权限',403)
  prior=c.execute(text('SELECT 1 FROM memberships WHERE warehouse_id=:w AND user_id=:u'),{'w':warehouse_id,'u':target}).first()
  if x.enabled and not prior:c.execute(text('INSERT INTO memberships(user_id,warehouse_id) VALUES(:u,:w)'),{'u':target,'w':warehouse_id})
  if not x.enabled and prior:c.execute(text('DELETE FROM memberships WHERE user_id=:u AND warehouse_id=:w'),{'u':target,'w':warehouse_id})
  audit(c,user['id'],warehouse_id,'分配员工仓库权限',json.dumps({'user_id':target,'enabled':x.enabled},ensure_ascii=False))
 return {'ok':True}


class CustomerSave(BaseModel):
 name:str=Field(min_length=1,max_length=120)
 contact:str=Field(default='',max_length=120)
 phone:str=Field(default='',max_length=80)
 username:str=Field(default='',max_length=100)
 password:str=Field(default='',max_length=200)
 warehouse_id:str=''

@app.get('/api/customer-management')
def customer_management(user=Depends(current)):
 permit(user,'customers')
 with engine.connect() as c:
  rows=[dict(r) for r in c.execute(text('SELECT id,name,contact,phone FROM customers WHERE deleted_at IS NULL ORDER BY name')).mappings()]
  for row in rows:
   row['accounts']=[] if user['role']!='老板' else [dict(a) for a in c.execute(text("SELECT u.id,u.username,u.active,u.client_warehouse_id warehouse_id,w.name warehouse_name FROM users u LEFT JOIN warehouses w ON w.id=u.client_warehouse_id WHERE u.role='客户' AND u.customer_id=:id ORDER BY u.username"),{'id':row['id']}).mappings()]
  return rows

def save_customer_record(c,x,user,cid=None):
 name=x.name.strip()
 if not name:fail('公司名称不能为空')
 if c.execute(text('SELECT id FROM customers WHERE LOWER(name)=LOWER(:n) AND id<>:id'),{'n':name,'id':cid or ''}).first():fail('客户公司已存在，请在列表为已有客户开通账号',409)
 if bool(x.username.strip())!=bool(x.password):fail('开通账号时必须同时填写账号和初始密码')
 if x.username.strip():
  allow(user,['老板'])
  if len(x.username.strip())<2 or len(x.password)<8:fail('账号至少2个字符，密码至少8位')
  access(c,user,x.warehouse_id)
  if c.execute(text('SELECT id FROM users WHERE username=:n'),{'n':x.username.strip()}).first():fail('账号已存在',409)
 if cid:
  if not c.execute(text('SELECT id FROM customers WHERE id=:id AND deleted_at IS NULL'),{'id':cid}).first():fail('客户不存在',404)
  c.execute(text('UPDATE customers SET name=:n,contact=:ct,phone=:p WHERE id=:id'),{'id':cid,'n':name,'ct':x.contact.strip(),'p':x.phone.strip()})
 else:
  cid=uid()
  c.execute(text('INSERT INTO customers(id,name,contact,phone) VALUES(:id,:n,:ct,:p)'),{'id':cid,'n':name,'ct':x.contact.strip(),'p':x.phone.strip()})
 if x.username.strip():
  try:c.execute(text("INSERT INTO users(id,username,password_hash,role,customer_id,client_warehouse_id) VALUES(:id,:u,:p,'客户',:cu,:w)"),{'id':uid(),'u':x.username.strip(),'p':pw_hash(x.password),'cu':cid,'w':x.warehouse_id})
  except IntegrityError:fail('账号已存在',409)
 audit(c,user['id'],None,'保存客户资料及账号',json.dumps({'customer_id':cid,'name':name,'username':x.username.strip()},ensure_ascii=False))
 return {'id':cid}

@app.post('/api/customer-management')
def create_customer_managed(x:CustomerSave,user=Depends(protect)):
 permit(user,'customers')
 try:
  with engine.begin() as c:return save_customer_record(c,x,user)
 except IntegrityError:fail('公司名称或账号已存在，请刷新后重试',409)

@app.patch('/api/customer-management/{cid}')
def edit_customer_managed(cid:str,x:CustomerSave,user=Depends(protect)):
 permit(user,'customers')
 try:
  with engine.begin() as c:return save_customer_record(c,x,user,cid)
 except IntegrityError:fail('公司名称或账号已存在，请刷新后重试',409)

class ClientPasswordReset(BaseModel):password:str=Field(min_length=8,max_length=200)
@app.post('/api/customer-accounts/{target}/password')
def reset_customer_password(target:str,x:ClientPasswordReset,user=Depends(protect)):
 allow(user,['老板'])
 with engine.begin() as c:
  if not c.execute(text("SELECT id FROM users WHERE id=:id AND role='客户'"),{'id':target}).first():fail('客户账号不存在',404)
  c.execute(text('UPDATE users SET password_hash=:p WHERE id=:id'),{'id':target,'p':pw_hash(x.password)})
  audit(c,user['id'],None,'重置客户密码',target)
 return {'ok':True}


class CustomerDelete(BaseModel):confirm_name:str=Field(min_length=1,max_length=120)
@app.delete('/api/customer-management/{cid}')
def delete_customer_managed(cid:str,x:CustomerDelete,user=Depends(protect)):
 allow(user,['老板','主管'])
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  suffix=' FOR UPDATE' if engine.dialect.name=='postgresql' else ''
  company=c.execute(text('SELECT id,name FROM customers WHERE id=:id AND deleted_at IS NULL'+suffix),{'id':cid}).mappings().first()
  if not company:fail('客户不存在或已删除',404)
  if x.confirm_name!=company['name']:fail('公司名称不一致，请输入完整公司名称确认',409)
  if c.execute(text('SELECT id FROM containers WHERE customer_id=:id LIMIT 1'),{'id':cid}).first() or c.execute(text('SELECT id FROM client_forecasts WHERE customer_id=:id LIMIT 1'),{'id':cid}).first():
   fail('该客户已有货柜、草稿或预报等业务记录，不能删除。请保留历史资料并停用客户账号',409)
  accounts=c.execute(text("SELECT id FROM users WHERE customer_id=:id AND role='客户'"),{'id':cid}).all()
  c.execute(text("UPDATE users SET active=0 WHERE customer_id=:id AND role='客户'"),{'id':cid})
  c.execute(text('UPDATE customers SET deleted_at=:t WHERE id=:id'),{'id':cid,'t':datetime.now(timezone.utc).isoformat()})
  audit(c,user['id'],None,'删除客户并注销账号',json.dumps({'customer_id':cid,'name':company['name'],'accounts':len(accounts)},ensure_ascii=False))
 return {'ok':True,'accounts_disabled':len(accounts)}


class ManualInboundPiece(BaseModel):
 weight:float|None=Field(default=None,ge=0,allow_inf_nan=False)
 length:float|None=Field(default=None,ge=0,allow_inf_nan=False)
 width:float|None=Field(default=None,ge=0,allow_inf_nan=False)
 height:float|None=Field(default=None,ge=0,allow_inf_nan=False)
 damage_status:str
class ManualInbound(BaseModel):
 request_id:str=Field(min_length=1,max_length=100)
 system_no:str=Field(min_length=1,max_length=150)
 main_barcode:str=Field(default='',max_length=200)
 transfer_no:str=Field(default='',max_length=200)
 channel:str=Field(default='',max_length=200)
 recipient:str=Field(default='',max_length=200)
 company:str=Field(default='',max_length=200)
 phone:str=Field(default='',max_length=100)
 email:str=Field(default='',max_length=200)
 country:str=Field(default='',max_length=100)
 city:str=Field(default='',max_length=150)
 postcode:str=Field(default='',max_length=50)
 address:str=Field(default='',max_length=2000)
 declared_value:float|None=Field(default=None,ge=0,allow_inf_nan=False)
 description:str=Field(default='',max_length=2000)
 remark:str=Field(default='',max_length=2000)
 pieces:list[ManualInboundPiece]=Field(min_length=1,max_length=100)

@app.post('/api/containers/{cid}/manual-inbound')
def manual_inbound(cid:str,x:ManualInbound,user=Depends(protect)):
 permit(user,'forecast')
 values=x.model_dump()
 for key,value in list(values.items()):
  if isinstance(value,str):values[key]=value.strip()
 if not values['system_no']:fail('请填写客户单号')
 if any(p.damage_status not in ('无破损','有破损') for p in x.pieces):fail('每件必须选择无破损或有破损')
 payload=json.dumps({'manual_inbound_container':cid,'data':values},sort_keys=True,ensure_ascii=False)
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  suffix=' FOR UPDATE' if engine.dialect.name=='postgresql' else ''
  ct=c.execute(text('SELECT * FROM containers WHERE id=:id AND deleted_at IS NULL'+suffix),{'id':cid}).mappings().first()
  if not ct:fail('货柜不存在',404)
  access(c,user,ct['warehouse_id'])
  prior=c.execute(text('SELECT * FROM operation_receipts WHERE request_id=:id'),{'id':x.request_id}).mappings().first()
  if prior:
   if prior['user_id']!=user['id'] or prior['payload']!=payload:fail('请求ID已用于其他操作',409)
   return dict(json.loads(prior['response']),duplicate_request=True)
  if ct['status']=='已关闭':fail('货柜已关闭，不能添加入库',409)
  if not c.execute(text('SELECT id FROM customers WHERE id=:id AND deleted_at IS NULL'),{'id':ct['customer_id']}).first():fail('当前货柜未关联有效客户，请先关联客户',409)
  if c.execute(text('SELECT id FROM orders WHERE container_id=:cid AND system_no=:no'),{'cid':cid,'no':values['system_no']}).first():fail('同柜客户单号已存在，请核对，不能重复添加',409)
  oid=uid();n=len(x.pieces);now=datetime.now(timezone.utc).isoformat()
  fields={k:v for k,v in values.items() if k not in ('request_id','pieces')}
  fields.update(id=oid,container_id=cid,expected=n,received=n,mode='待确认',goods_status='已入库',hold_label=int('扣货' in values['remark']),damage_label=int(any(p.damage_status=='有破损' for p in x.pieces)))
  try:c.execute(text('INSERT INTO orders('+','.join(fields)+') VALUES('+','.join(':'+k for k in fields)+')'),fields)
  except IntegrityError:fail('同柜客户单号已存在',409)
  for barcode in set(filter(None,[values['main_barcode'],values['transfer_no']])):
   c.execute(text('INSERT INTO barcodes(id,order_id,barcode) VALUES(:id,:o,:b)'),{'id':uid(),'o':oid,'b':barcode})
  for i,piece in enumerate(x.pieces,1):
   c.execute(text('INSERT INTO order_pieces(id,order_id,piece_no,weight,length,width,height,damage_status) VALUES(:id,:o,:n,:weight,:length,:width,:height,:damage_status)'),dict(piece.model_dump(),id=uid(),o=oid,n=i))
   c.execute(text("INSERT INTO piece_operations(id,request_id,order_id,warehouse_id,user_id,barcode,operation,result,reason) VALUES(:id,:req,:o,:w,:u,:b,'手动入库','成功',:reason)"),{'id':uid(),'req':uid(),'o':oid,'w':ct['warehouse_id'],'u':user['id'],'b':values['system_no'],'reason':f'第{i}件；{piece.damage_status}；手动添加并入库'})
  reason=f'手动添加并入库 {n}件'
  c.execute(text("INSERT INTO goods_events(id,order_id,kind,event_time,user_id,reason) VALUES(:id,:o,'入库',:t,:u,:r)"),{'id':uid(),'o':oid,'t':now,'u':user['id'],'r':reason})
  c.execute(text("INSERT INTO status_history(id,order_id,user_id,old_status,new_status,reason) VALUES(:id,:o,:u,'待入库','已入库',:r)"),{'id':uid(),'o':oid,'u':user['id'],'r':reason})
  audit(c,user['id'],ct['warehouse_id'],'手动添加并入库',json.dumps({'order_id':oid,'system_no':values['system_no'],'quantity':n,'request_id':x.request_id},ensure_ascii=False))
  result={'id':oid,'received':n,'expected':n,'status':'已入库','event_time':now}
  c.execute(text('INSERT INTO operation_receipts(request_id,user_id,payload,response) VALUES(:r,:u,:p,:s)'),{'r':x.request_id,'u':user['id'],'p':payload,'s':json.dumps(result,ensure_ascii=False)})
  return result


@app.get('/api/dashboard/details')
def dashboard_details(warehouse_id:str,kind:str,q:str='',customer_id:str='',container_id:str='',page:int=1,page_size:int=20,user=Depends(current)):
 permit(user,'read')
 if kind not in ('received','outbound','returns','containers','exceptions'):fail('不支持的统计类型')
 if page<1 or page_size<1 or page_size>100 or len(q)>200:fail('分页或搜索参数无效')
 with engine.connect() as c:
  access(c,user,warehouse_id)
  if kind=='exceptions':
   base="""SELECT e.id,e.container_id,ct.code container_code,ct.deleted_at,ct.customer_id,cu.name customer_name,e.barcode,COALESCE(o.system_no,e.barcode) system_no,o.main_barcode,o.id order_id,e.reason,e.created_at,e.status,1 quantity FROM exceptions e LEFT JOIN containers ct ON ct.id=e.container_id LEFT JOIN customers cu ON cu.id=ct.customer_id LEFT JOIN scans s ON s.id=e.scan_id LEFT JOIN orders o ON o.id=s.order_id WHERE e.warehouse_id=:w AND e.status='待处理'"""
  elif kind=='containers':
   base="""SELECT ct.id,ct.id container_id,ct.code container_code,ct.customer_id,cu.name customer_name,ct.status,COUNT(o.id) tickets,COALESCE(SUM(o.expected),0) expected,COALESCE(SUM(o.received),0) received,1 quantity FROM containers ct LEFT JOIN customers cu ON cu.id=ct.customer_id LEFT JOIN orders o ON o.container_id=ct.id WHERE ct.warehouse_id=:w AND ct.deleted_at IS NULL GROUP BY ct.id,ct.code,ct.customer_id,cu.name,ct.status"""
  else:
   received='CASE WHEN COALESCE(o.received,0)>=COALESCE(o.hold_count,0) AND COALESCE(o.received,0)>=0 THEN COALESCE(o.received,0) WHEN COALESCE(o.hold_count,0)>0 THEN o.hold_count ELSE 0 END'
   pending='CASE WHEN COALESCE(o.returned_count,0)>COALESCE(o.reout_count,0) THEN COALESCE(o.returned_count,0)-COALESCE(o.reout_count,0) ELSE 0 END'
   outside=f'CASE WHEN COALESCE(o.outbound_count,0)>({pending}) THEN COALESCE(o.outbound_count,0)-({pending}) ELSE 0 END'
   qty={'received':received,'outbound':outside,'returns':pending}[kind]
   base=f"SELECT o.id,o.id order_id,o.system_no,o.main_barcode,o.transfer_no,o.remark,o.hold_label,o.hold_count,o.expected,ct.id container_id,ct.code container_code,ct.customer_id,cu.name customer_name,({received}) received,({outside}) current_outbound,({pending}) pending_returns,({qty}) quantity FROM orders o JOIN containers ct ON ct.id=o.container_id LEFT JOIN customers cu ON cu.id=ct.customer_id WHERE ct.warehouse_id=:w AND ct.deleted_at IS NULL AND ({qty})>0"
  params={'w':warehouse_id,'cu':customer_id,'ct':container_id,'q':'%'+q.strip().lower().replace('!','!!').replace('%','!%').replace('_','!_')+'%'}
  where=" WHERE (:cu='' OR d.customer_id=:cu) AND (:ct='' OR d.container_id=:ct)"
  if q.strip():
   search=['container_code','customer_name']+([] if kind=='containers' else ['system_no','main_barcode'])+(['reason','barcode'] if kind=='exceptions' else ['remark','transfer_no'] if kind!='containers' else [])
   where+=' AND ('+' OR '.join(f"LOWER(COALESCE(d.{k},'')) LIKE :q ESCAPE '!'" for k in search)+')'
  source=' FROM ('+base+') d'
  summary=c.execute(text('SELECT COUNT(*) total,COALESCE(SUM(d.quantity),0) quantity'+source+where),params).mappings().first()
  total=int(summary['total']);pages=max(1,(total+page_size-1)//page_size);page=min(page,pages)
  sort='d.container_code,d.id' if kind=='containers' else 'd.created_at DESC,d.id' if kind=='exceptions' else 'd.container_code,d.system_no,d.id'
  rows=[dict(r) for r in c.execute(text('SELECT d.*'+source+where+' ORDER BY '+sort+' LIMIT :lim OFFSET :off'),dict(params,lim=page_size,off=(page-1)*page_size)).mappings()]
  options=[dict(r) for r in c.execute(text('SELECT DISTINCT d.customer_id,d.customer_name,d.container_id,d.container_code'+source+' ORDER BY d.container_code'),{'w':warehouse_id}).mappings()]
  return {'rows':rows,'total':total,'quantity':int(summary['quantity']),'page':page,'pages':pages,'page_size':page_size,'options':options}


def store_original_forecast(c,data,filename,user,customer_id=None,container_id=None):
 import base64
 name=filename.replace('\\','/').split('/')[-1][:255]
 if not name.lower().endswith(('.xls','.xlsx')):fail('请上传 xls 或 xlsx 原表')
 ident=uid()
 c.execute(text('INSERT INTO forecast_original_files(id,filename,content_base64,sha256,user_id,customer_id,container_id,created_at) VALUES(:id,:f,:data,:sha,:u,:cu,:ct,:t)'),{'id':ident,'f':name,'data':base64.b64encode(data).decode('ascii'),'sha':hashlib.sha256(data).hexdigest(),'u':user['id'],'cu':customer_id,'ct':container_id,'t':datetime.now(timezone.utc).isoformat()})
 return ident

def verify_original_source(c,source_id,user,customer_id=None):
 source=c.execute(text('SELECT id,user_id,customer_id,container_id FROM forecast_original_files WHERE id=:id'),{'id':source_id}).mappings().first()
 if not source:fail('原表记录不存在，请重新上传',409)
 if user['role']=='客户':
  if source['customer_id']!=customer_id:fail('无权使用其他公司的原表',403)
 elif source['user_id']!=user['id']:fail('无权使用他人上传的原表',403)
 return source

def bind_original_forecast(c,source_id,cid,user,customer_id=None):
 if not source_id:return
 source=verify_original_source(c,source_id,user,customer_id)
 if source['container_id'] and source['container_id']!=cid:fail('该原表已绑定其他货柜',409)
 c.execute(text('UPDATE forecast_original_files SET container_id=:ct WHERE id=:id'),{'ct':cid,'id':source_id})

def original_forecast_record(c,cid,user):
 ct=c.execute(text('SELECT * FROM containers WHERE id=:id AND deleted_at IS NULL'),{'id':cid}).mappings().first()
 if not ct:fail('货柜不存在',404)
 access(c,user,ct['warehouse_id'])
 source=c.execute(text('SELECT id,filename,sha256,created_at FROM forecast_original_files WHERE container_id=:id ORDER BY created_at DESC,id DESC LIMIT 1'),{'id':cid}).mappings().first()
 if not source:fail('该货柜未留存原始预报Excel，请补存原文件后再下载。系统当前数据无法原样恢复原文件格式。',404)
 return source

@app.get('/api/containers/{cid}/original-forecast/info')
def original_forecast_info(cid:str,user=Depends(current)):
 permit(user,'read')
 with engine.connect() as c:return dict(original_forecast_record(c,cid,user))

@app.get('/api/containers/{cid}/original-forecast')
def download_original_forecast(cid:str,user=Depends(current)):
 permit(user,'read')
 import base64
 from urllib.parse import quote
 with engine.connect() as c:
  source=original_forecast_record(c,cid,user)
  data=base64.b64decode(c.execute(text('SELECT content_base64 FROM forecast_original_files WHERE id=:id'),{'id':source['id']}).scalar())
 if hashlib.sha256(data).hexdigest()!=source['sha256']:fail('原表文件校验失败，请联系管理员',409)
 mime='application/vnd.ms-excel' if source['filename'].lower().endswith('.xls') else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
 return Response(data,media_type=mime,headers={'Content-Disposition':"attachment; filename*=UTF-8''"+quote(source['filename'],safe=''),'Cache-Control':'private, no-store','X-Content-Type-Options':'nosniff'})

@app.post('/api/containers/{cid}/original-forecast')
async def supplement_original_forecast(cid:str,file:UploadFile=File(...),user=Depends(protect)):
 permit(user,'forecast')
 with engine.connect() as c:
  ct=c.execute(text('SELECT * FROM containers WHERE id=:id AND deleted_at IS NULL'),{'id':cid}).mappings().first()
  if not ct:fail('货柜不存在',404)
  access(c,user,ct['warehouse_id'])
 data=await file.read(15_000_001)
 if len(data)>15_000_000:fail('Excel 超过15MB')
 if not (file.filename or '').lower().endswith(('.xls','.xlsx')):fail('请上传 xls 或 xlsx 文件')
 try:parse_excel(data,file.filename)
 except Exception:fail('文件无法识别为有效预报Excel')
 with engine.begin() as c:
  suffix=' FOR UPDATE' if engine.dialect.name=='postgresql' else ''
  ct=c.execute(text('SELECT * FROM containers WHERE id=:id AND deleted_at IS NULL'+suffix),{'id':cid}).mappings().first()
  if not ct:fail('货柜不存在',404)
  access(c,user,ct['warehouse_id'])
  if c.execute(text('SELECT id FROM forecast_original_files WHERE container_id=:id'),{'id':cid}).first():fail('该柜已经留存原表，补存不会覆盖原文件，请直接下载',409)
  ident=store_original_forecast(c,data,file.filename,user,ct['customer_id'],cid)
  audit(c,user['id'],ct['warehouse_id'],'补存原始预报表',json.dumps({'container_id':cid,'file_id':ident,'filename':file.filename},ensure_ascii=False))
 return {'ok':True}


def container_stock(c,cid):
 rows=c.execute(text('SELECT received,hold_count,outbound_count,returned_count,reout_count FROM orders WHERE container_id=:id'),{'id':cid}).mappings()
 return sum(max(0,max(0,int(o['received'] or 0),int(o['hold_count'] or 0))-max(0,int(o['outbound_count'] or 0)-max(0,int(o['returned_count'] or 0)-int(o['reout_count'] or 0)))) for o in rows)

@app.get('/api/containers/{cid}/force-delete-info')
def force_delete_info(cid:str,user=Depends(current)):
 permit(user,'force')
 with engine.connect() as c:
  ct=c.execute(text('SELECT * FROM containers WHERE id=:id AND deleted_at IS NULL'),{'id':cid}).mappings().first()
  if not ct:fail('货柜不存在或已作废',404)
  access(c,user,ct['warehouse_id'])
  return {'id':cid,'code':ct['code'],'in_stock':container_stock(c,cid)}

class ForceContainerDelete(BaseModel):
 code:str=Field(min_length=1,max_length=120)
 reason:str=Field(min_length=1,max_length=2000)
 expected_in_stock:int=Field(ge=0)
 request_id:str=Field(min_length=1,max_length=100)
@app.post('/api/containers/{cid}/force-delete')
def force_delete_container(cid:str,x:ForceContainerDelete,user=Depends(protect)):
 permit(user,'force')
 if not x.reason.strip():fail('请填写删除原因')
 payload=json.dumps({'force_delete_container':cid,**x.model_dump()},sort_keys=True,ensure_ascii=False)
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  suffix=' FOR UPDATE' if engine.dialect.name=='postgresql' else ''
  ct=c.execute(text('SELECT * FROM containers WHERE id=:id'+suffix),{'id':cid}).mappings().first()
  if not ct:fail('货柜不存在',404)
  access(c,user,ct['warehouse_id'])
  prior=c.execute(text('SELECT * FROM operation_receipts WHERE request_id=:r'),{'r':x.request_id}).mappings().first()
  if prior:
   if prior['user_id']!=user['id'] or prior['payload']!=payload:fail('请求ID已用于其他操作',409)
   return dict(json.loads(prior['response']),duplicate_request=True)
  if ct['deleted_at']:fail('货柜已作废，不能重复删除',409)
  if x.code!=ct['code']:fail('柜号不一致，请输入完整柜号确认',409)
  c.execute(text('SELECT id FROM orders WHERE container_id=:id'+suffix),{'id':cid}).all()
  stock=container_stock(c,cid)
  if stock!=x.expected_in_stock:fail('本柜库存发生变化，请关闭并重新打开删除窗口核对',409)
  stamp=datetime.now(timezone.utc).isoformat()
  c.execute(text("UPDATE containers SET deleted_at=:t,status='已作废' WHERE id=:id"),{'t':stamp,'id':cid})
  c.execute(text("UPDATE client_forecasts SET status='void',updated_at=:t,revision=revision+1 WHERE container_id=:id"),{'t':stamp,'id':cid})
  c.execute(text("UPDATE exceptions SET status='已作废' WHERE container_id=:id AND status='待处理'"),{'id':cid})
  audit(c,user['id'],ct['warehouse_id'],'强制删除货柜',json.dumps({'code':ct['code'],'in_stock':stock,'reason':x.reason.strip(),'history_preserved':True},ensure_ascii=False))
  result={'ok':True,'code':ct['code'],'removed_stock':stock}
  c.execute(text('INSERT INTO operation_receipts(request_id,user_id,payload,response) VALUES(:r,:u,:p,:s)'),{'r':x.request_id,'u':user['id'],'p':payload,'s':json.dumps(result,ensure_ascii=False)})
  return result

def readable_audit(c,action,raw,warehouse_id):
 import ast
 def resolve_order(ident):
  return c.execute(text('SELECT o.system_no FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.warehouse_id=:w'),{'id':ident,'w':warehouse_id}).scalar()
 def display_time(value):
  try:return datetime.fromisoformat(str(value).replace('Z','+00:00')).strftime('%Y年%m月%d日 %H:%M:%S')
  except (ValueError,TypeError):return str(value)
 labels={'order':'客户单号','system_no':'客户单号','piece_no':'件号','damage':'破损检查','damage_status':'破损状态','old':'原状态','new':'新状态','code':'柜号','orders':'预报票数','quantity':'件数','username':'账号','cycle':'退回轮次','repaired_at':'修复时间','reason':'原因','received':'实收件数','expected':'预报件数','before':'修改前','after':'修改后','name':'名称','enabled':'授权','role':'角色','active':'启用','weight':'重量kg','length':'长cm','width':'宽cm','height':'高cm','remark':'备注','channel':'渠道','country':'国家','city':'城市','postcode':'邮编','phone':'电话','email':'邮箱','recipient':'收件人','address':'地址','company':'公司','description':'物品描述','declared_value':'货值','filename':'原文件','main_barcode':'客户唛头','transfer_no':'新标签','in_stock':'移出统计的在库件数','damage_choices':'逐件破损检查','hold_confirmed':'扣货已确认','history_preserved':'历史记录保留','accounts':'注销账号数','quantity_removed':'移出件数','container_id':'货柜','customer_id':'客户公司','order_id':'客户单号','user_id':'操作账号'}
 hidden={'id','check_id','request_id','file_id','photo_id','payload','password_hash','source_id','sha256'}
 def value(k,v):
  if k in ('order_id','order'):return resolve_order(str(v)) or str(v)
  if k=='customer_id':return c.execute(text('SELECT name FROM customers WHERE id=:id'),{'id':v}).scalar() or '历史客户'
  if k=='container_id':return c.execute(text('SELECT code FROM containers WHERE id=:id AND warehouse_id=:w'),{'id':v,'w':warehouse_id}).scalar() or '历史货柜'
  if k=='user_id':return c.execute(text('SELECT username FROM users WHERE id=:id'),{'id':v}).scalar() or '历史账号'
  if k in ('repaired_at','event_time','created_at'):return display_time(v)
  if isinstance(v,dict):return describe(v)
  if isinstance(v,list):return '、'.join(str(x) for x in v[:30])
  if isinstance(v,bool):return '是' if v else '否'
  return str(v) if v is not None and v!='' else '未填写'
 def describe(obj):
  parts=[]
  for k,v in obj.items():
   if k in hidden:continue
   if k=='piece_no':parts.append('第'+str(v)+'件')
   elif k=='cycle':parts.append('第'+str(v)+'轮退回')
   elif k in ('order','system_no'):parts.append(value(k,v))
   else:parts.append(labels.get(k,k)+'：'+value(k,v))
  return ' · '.join(parts) or '操作已记录'
 raw=str(raw or '')
 try:obj=json.loads(raw)
 except (ValueError,TypeError):
  try:obj=ast.literal_eval(raw) if raw.startswith('{') else None
  except (ValueError,SyntaxError,MemoryError,RecursionError):obj=None
 if isinstance(obj,dict):return describe(obj)
 def replace_id(match):return resolve_order(match.group(0)) or ('照片记录' if '照片' in action else '历史记录')
 return re.sub(r'[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}',replace_id,raw) or '操作已记录'


def pc_order(c,oid,user,lock=False):
 suffix=' FOR UPDATE' if lock and engine.dialect.name=='postgresql' else ''
 if lock:c.execute(text('SELECT id FROM orders WHERE id=:id'+suffix),{'id':oid}).first()
 o=c.execute(text('SELECT o.*,ct.warehouse_id,ct.status container_status FROM orders o JOIN containers ct ON ct.id=o.container_id WHERE o.id=:id AND ct.deleted_at IS NULL'),{'id':oid}).mappings().first()
 if not o:fail('货物不存在或货柜已作废',404)
 access(c,user,o['warehouse_id'])
 return dict(o)

def pc_operation_info(c,o):
 keys=('received','outbound_count','return_target','returned_count','reout_count','return_cycle','hold_count','expected')
 snapshot={k:int(o[k] or 0) for k in keys}
 token=hashlib.sha256(json.dumps(snapshot,sort_keys=True).encode()).hexdigest()
 notes=[o['remark'] or '']+[r[0] for r in c.execute(text('SELECT note FROM goods_notes WHERE order_id=:id'),{'id':o['id']})]
 hold=bool(o['hold_label'] or any('扣货' in note for note in notes))
 legacy=bool(o['hold_prealert'] or o['hold_count'] or o['goods_status']=='扣件')
 expected=snapshot['expected'];received=snapshot['received'];out=snapshot['outbound_count'];target=snapshot['return_target'];ret=snapshot['returned_count'];reout=snapshot['reout_count']
 remaining={'入库':max(0,expected-received),'出库':max(0,received-out) if received>=expected and not target else 0,'退回':max(0,out if not target or reout==target else target-ret) if not(target and reout and reout<target) else 0,'重出':max(0,target-reout) if target and ret==target else 0}
 blocked=False
 try:ensure_damage_repaired(c,o)
 except HTTPException:blocked=True
 if blocked:remaining['出库']=remaining['重出']=0
 if legacy or o['container_status']=='已关闭':remaining={k:0 for k in remaining}
 return {'snapshot':token,'remaining':remaining,'hold':hold,'legacy_hold':legacy,'damage_blocked':blocked,'system_no':o['system_no'],'received':received,'expected':expected}

@app.get('/api/orders/{oid}/pc-operation-info')
def get_pc_operation_info(oid:str,user=Depends(current)):
 permit(user,'pc')
 with engine.connect() as c:return pc_operation_info(c,pc_order(c,oid,user))

class PCOperation(BaseModel):
 request_id:str=Field(min_length=1,max_length=100)
 operation:str
 quantity:int=Field(ge=1,le=1000)
 snapshot:str=Field(min_length=1,max_length=64)
 reason:str=Field(default='',max_length=2000)
 damage_choices:list[str]=Field(default_factory=list,max_length=1000)
 hold_confirmed:bool=False

@app.post('/api/orders/{oid}/pc-operation')
def pc_operation(oid:str,x:PCOperation,user=Depends(protect)):
 permit(user,'pc')
 if x.operation not in ('入库','出库','退回','重出'):fail('不支持的PC操作')
 payload=json.dumps({'pc_order':oid,**x.model_dump()},sort_keys=True,ensure_ascii=False)
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  o=pc_order(c,oid,user,True)
  prior=c.execute(text('SELECT * FROM operation_receipts WHERE request_id=:r'),{'r':x.request_id}).mappings().first()
  if prior:
   if prior['user_id']!=user['id'] or prior['payload']!=payload:fail('请求ID已用于其他操作',409)
   return dict(json.loads(prior['response']),duplicate_request=True)
  info=pc_operation_info(c,o);op=x.operation;n=x.quantity
  if info['snapshot']!=x.snapshot:fail('货物数量已变化，请刷新操作区重新核对',409)
  if o['container_status']=='已关闭':fail('货柜已关闭',409)
  if info['legacy_hold']:fail('存在旧版扣件流程，请先由主管核对并解除旧版扣件后操作',409)
  if info['hold'] and not x.hold_confirmed:fail('此票为扣货，必须确认指定区域和本次操作',409)
  if op in ('出库','重出'):ensure_damage_repaired(c,o)
  if n>info['remaining'][op]:fail('操作数量超过当前可操作件数，或尚未满足当前流程条件',409)
  if op in ('入库','退回') and (len(x.damage_choices)!=n or any(v not in ('无破损','有破损') for v in x.damage_choices)):fail('必须为本次每件货物选择破损情况')
  reason=x.reason.strip() or 'PC直接操作';stamp=datetime.now(timezone.utc).isoformat();old=o['goods_status'];new=old
  if op=='入库':
   for i,damage in enumerate(x.damage_choices,1):
    piece_no=int(o['received'] or 0)+i
    existing=c.execute(text('SELECT id FROM order_pieces WHERE order_id=:o AND piece_no=:n'),{'o':oid,'n':piece_no}).first()
    if existing:c.execute(text('UPDATE order_pieces SET damage_status=:d WHERE order_id=:o AND piece_no=:n'),{'d':damage,'o':oid,'n':piece_no})
    else:c.execute(text('INSERT INTO order_pieces(id,order_id,piece_no,damage_status) VALUES(:id,:o,:n,:d)'),{'id':uid(),'o':oid,'n':piece_no,'d':damage})
   c.execute(text('UPDATE orders SET received=received+:n WHERE id=:id'),{'n':n,'id':oid})
   if int(o['received'] or 0)+n==o['expected']:new='已入库'
  elif op=='出库':
   c.execute(text('UPDATE orders SET outbound_count=outbound_count+:n WHERE id=:id'),{'n':n,'id':oid})
   if int(o['outbound_count'] or 0)+n==o['expected']:new='已出库'
  elif op=='退回':
   target=int(o['return_target'] or 0);ret=int(o['returned_count'] or 0);cycle=int(o['return_cycle'] or 0)
   if not target or int(o['reout_count'] or 0)==target:target=int(o['outbound_count'] or 0);ret=0;cycle+=1
   for damage in x.damage_choices:c.execute(text('INSERT INTO return_checks(id,order_id,cycle,request_id,quantity,scope,damage_choice,user_id,reason) VALUES(:id,:o,:cy,:req,1,\'single\',:d,:u,:r)'),{'id':uid(),'o':oid,'cy':cycle,'req':uid(),'d':damage,'u':user['id'],'r':'PC退回；'+reason})
   c.execute(text('UPDATE orders SET return_target=:t,returned_count=:ret,reout_count=0,return_cycle=:cy WHERE id=:id'),{'t':target,'ret':ret+n,'cy':cycle,'id':oid});new='退回'
  else:
   c.execute(text('UPDATE orders SET reout_count=reout_count+:n WHERE id=:id'),{'n':n,'id':oid});new='重出'
  if '有破损' in x.damage_choices:c.execute(text('UPDATE orders SET damage_label=1 WHERE id=:id'),{'id':oid})
  c.execute(text('UPDATE orders SET goods_status=:s WHERE id=:id'),{'s':new,'id':oid})
  if new!=old:c.execute(text('INSERT INTO status_history(id,order_id,user_id,old_status,new_status,reason) VALUES(:id,:o,:u,:old,:new,:r)'),{'id':uid(),'o':oid,'u':user['id'],'old':old,'new':new,'r':f'PC{op} {n}件；'+reason})
  for i in range(n):c.execute(text('INSERT INTO piece_operations(id,request_id,order_id,warehouse_id,user_id,barcode,operation,result,reason) VALUES(:id,:req,:o,:w,:u,:b,:op,\'成功\',:r)'),{'id':uid(),'req':uid(),'o':oid,'w':o['warehouse_id'],'u':user['id'],'b':o['system_no'],'op':op,'r':f'PC直接{op} 第{i+1}/{n}件；'+(x.damage_choices[i]+'；' if op in ('入库','退回') else '')+reason})
  c.execute(text('INSERT INTO goods_events(id,order_id,kind,event_time,user_id,reason) VALUES(:id,:o,:k,:t,:u,:r)'),{'id':uid(),'o':oid,'k':op,'t':stamp,'u':user['id'],'r':f'PC直接操作 {n}件；'+reason})
  audit(c,user['id'],o['warehouse_id'],'PC直接'+op,json.dumps({'order':o['system_no'],'quantity':n,'reason':reason,'damage_choices':x.damage_choices,'hold_confirmed':x.hold_confirmed},ensure_ascii=False))
  result={'ok':True,'operation':op,'quantity':n,'order_id':oid}
  c.execute(text('INSERT INTO operation_receipts(request_id,user_id,payload,response) VALUES(:r,:u,:p,:s)'),{'r':x.request_id,'u':user['id'],'p':payload,'s':json.dumps(result,ensure_ascii=False)})
  return result


PERMISSION_LABELS={'pc':'PC入库 / 出库 / 退回 / 重出','pda':'PDA扫码作业','repair':'确认破损修复','forecast':'预报管理','edit':'编辑货物资料','notes':'备注与照片','force':'强制删除货柜','customers':'客户资料管理','warehouses':'仓库管理','staff':'员工账号管理','read':'货物及统计查询'}
ROLE_PERMISSIONS={
 '老板':set(PERMISSION_LABELS),'主管':set(PERMISSION_LABELS),
 '客服':{'forecast','edit','notes','customers','read'},
 '仓库管理员':{'pc','pda','edit','notes','read'},
 'PDA操作员':{'pda','notes','read'},'PDA员工':{'pda','notes','read'},'只读人员':{'read'},'客户':set()}
def effective_permissions(user,c=None):
 if user['role']=='客户':return set()
 defaults=set(ROLE_PERMISSIONS.get(user['role'],set()))
 if user['role']=='老板':return defaults
 def read(conn):return conn.execute(text('SELECT permission,enabled FROM employee_permissions WHERE user_id=:id'),{'id':user['id']}).all()
 if c is None:
  with engine.connect() as conn:rows=read(conn)
 else:rows=read(c)
 for key,enabled in rows:
  if key not in PERMISSION_LABELS:continue
  if enabled:defaults.add(key)
  else:defaults.discard(key)
 return defaults
def permit(user,key):
 if key not in effective_permissions(user):fail('未获授权：'+PERMISSION_LABELS[key],403)

class StaffPermissionsEdit(BaseModel):
 permissions:list[str]=Field(default_factory=list,max_length=20)
 revision:int=Field(ge=0)
 restore_defaults:bool=False

def permission_target(c,target,user):
 allow(user,['老板','主管']);permit(user,'staff')
 employee=c.execute(text('SELECT id,username,role FROM users WHERE id=:id'),{'id':target}).mappings().first()
 if not employee or employee['role']=='客户':fail('请选择公司员工，客户账号独立管理',404)
 if target==user['id']:fail('请由其他老板或主管调整当前账号的权限',403)
 if employee['role']=='老板':fail('老板权限不可由此界面修改',403)
 # A manager must be authorized for every warehouse assigned to the target.
 for row in c.execute(text('SELECT warehouse_id FROM memberships WHERE user_id=:id'),{'id':target}):access(c,user,row[0])
 return employee

def staff_permission_data(c,employee):
 revision=c.execute(text('SELECT revision FROM employee_permission_versions WHERE user_id=:id'),{'id':employee['id']}).scalar() or 0
 custom=c.execute(text('SELECT COUNT(*) FROM employee_permissions WHERE user_id=:id'),{'id':employee['id']}).scalar()>0
 warehouses=[dict(r) for r in c.execute(text('SELECT w.id,w.name FROM warehouses w JOIN memberships m ON m.warehouse_id=w.id WHERE m.user_id=:id ORDER BY w.name'),{'id':employee['id']}).mappings()]
 return {'employee':dict(employee),'effective':sorted(effective_permissions(employee,c)),'defaults':sorted(ROLE_PERMISSIONS.get(employee['role'],set())),'labels':PERMISSION_LABELS,'revision':int(revision),'custom':custom,'warehouses':warehouses}

@app.get('/api/users/{target}/permissions')
def get_staff_permissions(target:str,user=Depends(current)):
 with engine.connect() as c:return staff_permission_data(c,permission_target(c,target,user))

@app.put('/api/users/{target}/permissions')
def save_staff_permissions(target:str,x:StaffPermissionsEdit,user=Depends(protect)):
 if set(x.permissions)-set(PERMISSION_LABELS):fail('存在不支持的权限')
 with engine.begin() as c:
  if engine.dialect.name=='sqlite':c.exec_driver_sql('BEGIN IMMEDIATE')
  suffix=' FOR UPDATE' if engine.dialect.name=='postgresql' else ''
  c.execute(text('SELECT id FROM users WHERE id=:id'+suffix),{'id':target}).first()
  employee=permission_target(c,target,user)
  prior=staff_permission_data(c,employee)
  if prior['revision']!=x.revision:fail('权限已被其他管理员修改，请重新打开核对',409)
  selected=set(ROLE_PERMISSIONS.get(employee['role'],set())) if x.restore_defaults else set(x.permissions)
  if selected-{'read'} and 'read' not in selected:fail('分配业务功能时请同时勾选货物及统计查询')
  c.execute(text('DELETE FROM employee_permissions WHERE user_id=:id'),{'id':target})
  if not x.restore_defaults:
   for key in PERMISSION_LABELS:c.execute(text('INSERT INTO employee_permissions(user_id,permission,enabled) VALUES(:id,:p,:v)'),{'id':target,'p':key,'v':int(key in selected)})
  if prior['revision']:c.execute(text('UPDATE employee_permission_versions SET revision=revision+1 WHERE user_id=:id'),{'id':target})
  else:c.execute(text('INSERT INTO employee_permission_versions(user_id,revision) VALUES(:id,1)'),{'id':target})
  audit(c,user['id'],None,'设置员工权限',json.dumps({'username':employee['username'],'before':[PERMISSION_LABELS[p] for p in prior['effective']],'after':[PERMISSION_LABELS[p] for p in sorted(selected)]},ensure_ascii=False))
  return staff_permission_data(c,employee)
